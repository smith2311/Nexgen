"""Excel export of the Vendor x RFI price-premium heatmap.

Takes the dict :func:`vendorlink.quoteheatmap.build` already produced for the
screen and writes it to a workbook. It is deliberately fed that dict rather
than the bids, so the file cannot disagree with the matrix the user was looking
at when they clicked Export — same filters, same dropped columns, same numbers.

Four sheets, because the on-screen grid alone does not survive being pulled out
of its page:

``Heatmap``
    The matrix as plotted: vendors down, RFI columns across, each cell the
    premium above that column's cheapest bid, filled with the same ramp colour
    the screen used. This is the picture.

``RFI columns``
    What each column actually is. The headers on screen are truncated to 34
    characters and carry their detail in a tooltip, which does not export.

``Cell data``
    The same numbers in tidy/long form -- one row per vendor-RFI-currency, with
    the vendor's total, the column's cheapest, the premium and its band. This
    is the sheet to pivot or re-plot from; the matrix is not a data table.

``Notes``
    The caveats. The heatmap is built to avoid several specific ways of
    misreading bid data -- premiums are not prices, currencies are never mixed,
    single-bidder columns are dropped -- and a spreadsheet that carried the
    numbers without those rules would reintroduce every one of them the moment
    someone sorted a column.

No Django imports: the workbook is a pure function of the heat dict, so it can
be tested without a database.
"""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .quoteheatmap import PREMIUM_BANDS

#: House header style, matching the year-end backup workbook.
_HEADER_FILL = PatternFill("solid", fgColor="0076C8")
_HEADER_FONT = Font(bold=True, color="FFFFFF")

#: Premium cells are numbers, not the "+40%" strings the screen shows. A text
#: cell cannot be sorted, averaged or charted, which is most of the reason to
#: want this in Excel at all.
_PCT_FORMAT = '0.0"%"'
_MONEY_FORMAT = "#,##0.00"


def _text(value):
    return "" if value is None else str(value)


def _autosize(sheet, min_width=10, max_width=48):
    for column in sheet.columns:
        longest = max((len(_text(cell.value)) for cell in column), default=0)
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(max(longest + 2, min_width), max_width)


def _style_header(sheet, row=1):
    for cell in sheet[row]:
        if cell.value is not None:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _column_label(meta, references):
    """Header for one matrix column: currency, reference and subject.

    The currency leads because it is the one thing that makes two columns of
    the same RFI different, and the whole premium is computed within it.
    """
    reference = references.get(meta["request_pk"]) if references else None
    parts = [meta["currency"]]
    if reference:
        parts.append(reference)
    parts.append(meta["subject"])
    return " · ".join(part for part in parts if part)


def _sheet_heatmap(workbook, heat, references):
    sheet = workbook.create_sheet("Heatmap")
    columns = heat["columns"]

    header = ["Vendor"] + [_column_label(meta, references) for meta in columns]
    header += ["Bids", "Wins", "Avg premium %"]
    sheet.append(header)
    _style_header(sheet)
    sheet.row_dimensions[1].height = 46

    for row in heat["rows"]:
        sheet.append([row["vendor"]])
        excel_row = sheet.max_row
        for index, cell in enumerate(row["cells"], start=2):
            target = sheet.cell(row=excel_row, column=index)
            if not cell.get("has_bid"):
                # Left empty on purpose. A 0 here would read as "cheapest",
                # when it means the vendor did not bid on this RFI at all --
                # the single most dangerous thing this export could get wrong.
                continue
            target.value = round(cell["premium"], 1)
            target.number_format = _PCT_FORMAT
            target.fill = PatternFill("solid", fgColor=cell["bg"].lstrip("#"))
            target.font = Font(color=cell["ink"].lstrip("#"), bold=bool(cell["is_best"]))
            target.alignment = Alignment(horizontal="center")

        summary_start = len(columns) + 2
        sheet.cell(row=excel_row, column=summary_start, value=row["bids"])
        sheet.cell(row=excel_row, column=summary_start + 1, value=row["wins"])
        average = sheet.cell(row=excel_row, column=summary_start + 2, value=row["avg_premium"])
        if row["avg_premium"] is not None:
            average.number_format = _PCT_FORMAT

    sheet.freeze_panes = "B2"
    _autosize(sheet, min_width=12, max_width=34)
    sheet.column_dimensions["A"].width = 34
    return sheet


def _sheet_columns(workbook, heat, references):
    sheet = workbook.create_sheet("RFI columns")
    sheet.append([
        "Column", "Reference", "Subject", "Country", "Currency",
        "Bidders", "Cheapest total", "Raised", "Split across currencies",
    ])
    _style_header(sheet)

    for position, meta in enumerate(heat["columns"], start=1):
        created = meta.get("created")
        sheet.append([
            position,
            references.get(meta["request_pk"], "") if references else "",
            meta["subject"],
            meta["country"],
            meta["currency"],
            meta["bidders"],
            meta["cheapest"],
            # Excel cannot hold a tz-aware datetime.
            created.replace(tzinfo=None) if created is not None else None,
            "Yes" if meta["multi_currency"] else "No",
        ])
        sheet.cell(row=sheet.max_row, column=7).number_format = _MONEY_FORMAT

    sheet.freeze_panes = "A2"
    _autosize(sheet)
    return sheet


def _sheet_cells(workbook, heat, references):
    sheet = workbook.create_sheet("Cell data")
    sheet.append([
        "Vendor", "Reference", "RFI", "Country", "Currency",
        "Vendor total", "Cheapest total", "Premium %", "Band", "Is cheapest",
    ])
    _style_header(sheet)

    by_position = {index: meta for index, meta in enumerate(heat["columns"])}
    for row in heat["rows"]:
        for index, cell in enumerate(row["cells"]):
            if not cell.get("has_bid"):
                continue
            meta = by_position[index]
            sheet.append([
                row["vendor"],
                references.get(meta["request_pk"], "") if references else "",
                meta["subject"],
                meta["country"],
                cell["currency"],
                cell["total"],
                meta["cheapest"],
                round(cell["premium"], 1),
                cell["band"],
                "Yes" if cell["is_best"] else "No",
            ])
            sheet.cell(row=sheet.max_row, column=6).number_format = _MONEY_FORMAT
            sheet.cell(row=sheet.max_row, column=7).number_format = _MONEY_FORMAT
            sheet.cell(row=sheet.max_row, column=8).number_format = _PCT_FORMAT

    sheet.freeze_panes = "A2"
    _autosize(sheet)
    return sheet


def _sheet_notes(workbook, heat, filters=None):
    sheet = workbook.create_sheet("Notes")
    sheet.append(["Field", "Value"])
    _style_header(sheet)

    notes = [
        ("Vendors", heat["vendor_count"]),
        ("Columns", heat["column_count"]),
        ("Columns dropped (single bidder)", heat["dropped_single"]),
        ("Columns dropped (over the display cap)", heat["truncated"]),
        ("RFIs split across currencies", heat["split_requests"]),
    ]
    if filters:
        for label, value in filters.items():
            notes.append((f"Filter — {label}", value or "(all)"))
    for label, value in notes:
        sheet.append([label, value])

    sheet.append([])
    sheet.append(["How to read this"])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

    for line in (
        "Cells are premiums, not prices: each is the percentage above the "
        "cheapest bid for that column. 0% is the winner.",

        "Raw totals are NOT comparable across RFIs -- a national warehousing "
        "deal and a single-lane courier quote differ by orders of magnitude. "
        "That is why the matrix holds premiums.",

        "Currencies are never mixed. A premium is computed strictly within one "
        "currency for one RFI, so an RFI that drew bids in two currencies "
        "becomes two columns. There is no exchange-rate table behind this "
        "file; do not compare a EUR column against a USD one.",

        "An empty cell means the vendor did not bid on that RFI. It does not "
        "mean zero, and it does not mean cheapest.",

        "Columns with only one bidder are excluded: a premium measured against "
        "yourself is always 0% and would paint a misleadingly perfect cell.",

        "Avg premium % is averaged over the RFIs each vendor actually bid on, "
        "not over every column, so a vendor who bid twice is not penalised for "
        "the RFIs they were never invited to.",

        "Bids with no costs filled in are skipped rather than counted as a "
        "quote of zero, which would invent a winner undercutting everyone.",
    ):
        sheet.append([line])
        sheet.cell(row=sheet.max_row, column=1).alignment = Alignment(wrap_text=True)

    sheet.append([])
    sheet.append(["Premium bands"])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    for index, step in enumerate(heat["steps"]):
        sheet.append([step["label"]])
        cell = sheet.cell(row=sheet.max_row, column=1)
        cell.fill = PatternFill("solid", fgColor=step["bg"].lstrip("#"))
        cell.font = Font(color=step["ink"].lstrip("#"))
        # Guards the ramp against drifting out of step with the band labels.
        assert step["label"] == PREMIUM_BANDS[index][1]

    sheet.column_dimensions["A"].width = 96
    sheet.column_dimensions["B"].width = 18
    return sheet


def build_workbook(heat, references=None, filters=None):
    """Workbook for a built heat dict. ``references`` maps request pk -> RFI id."""
    references = references or {}
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    if not heat.get("has_data"):
        sheet = workbook.create_sheet("Heatmap")
        sheet.append(["No comparable bids for the current filters."])
        sheet.append([
            "The heatmap needs at least one RFI with two or more vendors "
            "quoting in the same currency."
        ])
        _sheet_notes(workbook, heat, filters)
        return workbook

    _sheet_heatmap(workbook, heat, references)
    _sheet_columns(workbook, heat, references)
    _sheet_cells(workbook, heat, references)
    _sheet_notes(workbook, heat, filters)
    return workbook


def workbook_bytes(heat, references=None, filters=None):
    """The workbook as bytes, ready for an HttpResponse or a file."""
    output = io.BytesIO()
    build_workbook(heat, references, filters).save(output)
    return output.getvalue()
