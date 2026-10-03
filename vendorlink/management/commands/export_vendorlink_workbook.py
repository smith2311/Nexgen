"""Export the CURRENT VendorLink database state to an Excel workbook.

Nothing here is hardcoded — every row is read live from the database, so the
workbook always reflects real data. Run the seeder first if the DB is empty:

    python manage.py seed_vendorlink_demo --reset
    python manage.py export_vendorlink_workbook --out /app/media/vendorlink_data.xlsx

The AI's verdict is deliberately NOT exported. Which vendor offers the best
price for a request is decided by the evaluator at review time, from the request
page — it is not a spreadsheet column.
"""

from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.core.management.base import BaseCommand
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from vendorlink.capabilities import capability_rows
from vendorlink.models import Vendor, VendorBid, VendorRegistration, VendorRequest, VendorScore

BLUE, NAVY = "0076C8", "0F1F3A"
GREY, GREEN, AMBER, RED = "F1F5F9", "DCFCE7", "FEF3C7", "FEE2E2"

HDR_FILL = PatternFill("solid", fgColor=BLUE)
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
_thin = Side(style="thin", color="D9E2EC")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

COST_FIELDS = ["storage_cost", "inbound_cost", "outbound_cost", "pick_cost",
               "packing_cost", "last_mile_charges", "manpower_charges"]
# Read at call time, not import: the capability master is a query, and this
# command is loaded by manage.py before the database is necessarily reachable.


def cap_columns():
    """[(code, label), …] for the active capabilities, in master order."""
    return [(code, label) for code, label, _icon in capability_rows()]


def yn(v):
    return "Yes" if v else "No"


def bid_total(b):
    return sum(getattr(b, f) or Decimal(0) for f in COST_FIELDS)


class Command(BaseCommand):
    help = "Export live VendorLink data (registrations, vendors, scores, requests, bids) to Excel."

    def add_arguments(self, parser):
        parser.add_argument("--out", default="/app/media/vendorlink_data.xlsx",
                           help="Output .xlsx path.")

    # ── styling helpers ───────────────────────────────────────────────────
    def _head(self, ws, headers, row=1):
        ws.append(headers) if row == ws.max_row + 1 or ws.max_row < row else None
        for c, _ in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=c)
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30
        # String ref: ws.cell() would create the row and shift later appends down.
        ws.freeze_panes = f"A{row + 1}"

    def _finish(self, ws, first_row, ncols, widths):
        for r in range(first_row, ws.max_row + 1):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BOX
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            if (r - first_row) % 2:
                for c in range(1, ncols + 1):
                    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FAFBFD")
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _pill(self, ws, col, first_row, mapping, default=GREY):
        for r in range(first_row, ws.max_row + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = PatternFill("solid", fgColor=mapping.get(cell.value, default))
            cell.font = Font(bold=True, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # ──────────────────────────────────────────────────────────────────────
    def handle(self, *args, **opts):
        wb = openpyxl.Workbook()
        self.sheet_flow(wb)
        n_reg = self.sheet_registrations(wb)
        n_ven = self.sheet_vendor_master(wb)
        n_acc = self.sheet_accounts(wb)
        n_sco = self.sheet_scores(wb)
        n_req = self.sheet_requests(wb)
        n_bid = self.sheet_bids(wb)
        self.sheet_dictionary(wb)

        out = opts["out"]
        wb.save(out)
        self.stdout.write(self.style.SUCCESS(f"written: {out}"))
        self.stdout.write(
            f"  {n_reg} registrations, {n_ven} vendors, {n_acc} portal accounts, "
            f"{n_sco} scores, {n_req} requests, {n_bid} bids")
        self.stdout.write(self.style.WARNING(
            "  AI verdict intentionally not exported — run the evaluator on a request "
            "and read the recommendation, reasons and points in the app."))

    # ── 0. Process Flow (documentation, not data) ─────────────────────────
    def sheet_flow(self, wb):
        ws = wb.active
        ws.title = "0. Process Flow"
        ws["A1"] = "Meril VendorLink — Implemented Process Flow"
        ws["A1"].font = Font(bold=True, size=15, color=NAVY)
        ws.merge_cells("A1:E1")
        ws["A2"] = ("Every data sheet in this file is exported LIVE from the database. "
                    "Test inboxes are real — approving a vendor genuinely sends mail.")
        ws["A2"].font = Font(italic=True, size=10, color="64748B")
        ws.merge_cells("A2:E2")
        ws.append([])
        ws.append(["Stage", "What happens", "Actor", "Where in the app", "Sheet"])
        self._head(ws, ["Stage", "What happens", "Actor", "Where in the app", "Sheet"], row=4)
        rows = [
            ("1. Self-registration",
             "3PL submits the public form. Email verified by OTP; math CAPTCHA blocks bots; "
             "duplicate email refused via modal.",
             "Vendor (public)", "/vendorlink/register/", "1. Registrations"),
            ("2. Admin review", "Global SCM approves or rejects, with an optional reason "
             "emailed to the vendor.", "VendorLink Admin", "/vendorlink/registrations/",
             "1. Registrations"),
            ("3. Approval effects",
             "Creates the Vendor Master record, generates an 8-char temp password, creates a "
             "user in group VendorUser, emails credentials.",
             "System", "approve_registration()", "2. Vendor Master / 3. Portal Accounts"),
            ("4. First login", "Middleware forces a password change before any other page "
             "loads.", "Vendor", "/vendorlink/vendor/change-password/", "3. Portal Accounts"),
            ("5. Vendor scoring", "Admin rates facility, support, compliance, delivery 1-5 "
             "and may blacklist. Feeds the evaluator.", "VendorLink Admin",
             "Vendor detail page", "4. Vendor Scores"),
            ("6. Raise a request", "Broadcast a structured RFI to every active vendor in a "
             "country, with a bid window.", "Staff", "/vendorlink/requests/new/",
             "5. Requests"),
            ("7. Vendor bids", "Vendor quotes across 7 cost lines. Revisions allowed; staff "
             "may flag one as selected.", "Vendor", "Vendor portal", "6. Bids"),
            ("8. AI evaluation",
             "OpenAI infers the required services, marks vendors ineligible if blacklisted or "
             "missing one, then recommends the cheapest ELIGIBLE vendor with reasons and "
             "points. NOT exported to Excel — read it in the app.",
             "System (OpenAI)", "/vendorlink/requests/<id>/evaluate/", "— in-app only —"),
        ]
        for r in rows:
            ws.append(list(r))
        self._finish(ws, 5, 5, [22, 80, 18, 34, 32])

    # ── 1. Registrations ─────────────────────────────────────────────────
    def sheet_registrations(self, wb):
        ws = wb.create_sheet("1. Registrations")
        caps = cap_columns()
        hdr = (["SR#", "Company Name", "Country", "State", "City", "Warehouse Location",
                "Contact Person", "Email", "Mobile", "Website"]
               + [label for _code, label in caps] +
               ["Status", "Submitted At", "Reviewed By", "Reviewed At", "Rejection Reason"])
        ws.append(hdr)
        self._head(ws, hdr)
        qs = (
            VendorRegistration.objects.select_related("reviewed_by")
            .prefetch_related("capabilities")
            .order_by("status", "vendor_name")
        )
        for i, r in enumerate(qs, start=1):
            rb = r.reviewed_by
            ws.append([i, r.vendor_name, r.country, r.state, r.city, r.warehouse_location,
                       r.contact_person, r.email, r.mobile, r.website] +
                      [yn(r.has_capability(code)) for code, _label in caps] +
                      [r.status,
                       r.submitted_at.strftime("%Y-%m-%d %H:%M") if r.submitted_at else "",
                       (rb.get_full_name() or rb.username) if rb else "",
                       r.reviewed_at.strftime("%Y-%m-%d %H:%M") if r.reviewed_at else "",
                       r.rejection_reason])
        self._finish(ws, 2, len(hdr),
                     [6, 32, 20, 18, 16, 40, 20, 38, 20, 30] + [13] * len(caps) + [12, 18, 18, 18, 55])
        self._pill(ws, hdr.index("Status") + 1, 2,
                   {"approved": GREEN, "pending": AMBER, "rejected": RED})
        return ws.max_row - 1

    # ── 2. Vendor Master (matches the app's own import format) ────────────
    def sheet_vendor_master(self, wb):
        ws = wb.create_sheet("2. Vendor Master")
        caps = cap_columns()
        ws["A1"] = ("IMPORT-READY — these columns match the app's export/import format exactly "
                    "(Vendor Master > Import). Delete this banner row before importing.")
        ws["A1"].font = Font(bold=True, size=10, color="92400E")
        ws["A1"].fill = PatternFill("solid", fgColor=AMBER)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10 + len(caps))
        ws.row_dimensions[1].height = 26

        hdr = (["SR#", "Vendor Name", "Country", "State", "City", "Warehouse Location", "Contact Person",
                "Email", "Mobile", "Website"] + [label for _code, label in caps])
        ws.append(hdr)
        self._head(ws, hdr, row=2)
        vendor_qs = (
            Vendor.objects.filter(is_active=True)
            .prefetch_related("capabilities")
            .order_by("vendor_name")
        )
        for i, v in enumerate(vendor_qs, 1):
            ws.append([i, v.vendor_name, v.country, v.state, v.city, v.warehouse_location, v.contact_person,
                       v.email, v.mobile, v.website] +
                      [yn(v.has_capability(code)) for code, _label in caps])
        self._finish(ws, 3, len(hdr), [6, 34, 22, 18, 16, 42, 20, 38, 20, 30] + [13] * len(caps))
        return ws.max_row - 2

    # ── 3. Portal Accounts ───────────────────────────────────────────────
    def sheet_accounts(self, wb):
        ws = wb.create_sheet("3. Portal Accounts")
        ws["A1"] = ("Temp passwords are NOT exported — they are hashed on save and emailed "
                    "once. The seeder prints them: manage.py seed_vendorlink_demo")
        ws["A1"].font = Font(bold=True, size=10, color="92400E")
        ws["A1"].fill = PatternFill("solid", fgColor=AMBER)
        ws.merge_cells("A1:G1")
        ws.row_dimensions[1].height = 26

        hdr = ["SR#", "Username (= email)", "Vendor", "Groups", "Active",
               "Must Change Password", "Last Login"]
        ws.append(hdr)
        self._head(ws, hdr, row=2)
        users = User.objects.filter(groups__name="VendorUser").distinct().order_by("username")
        for i, u in enumerate(users, start=1):
            reg = VendorRegistration.objects.filter(email=u.email).first()
            try:
                must = yn(u.profile.force_password_change)
            except Exception:
                must = "—"
            ws.append([i, u.username, reg.vendor_name if reg else "—",
                       ", ".join(g.name for g in u.groups.all()), yn(u.is_active), must,
                       u.last_login.strftime("%Y-%m-%d %H:%M") if u.last_login
                       else "never logged in"])
        self._finish(ws, 3, len(hdr), [6, 40, 34, 16, 10, 22, 20])
        self._pill(ws, hdr.index("Must Change Password") + 1, 3, {"Yes": AMBER, "No": GREEN})
        return ws.max_row - 2

    # ── 4. Vendor Scores ─────────────────────────────────────────────────
    def sheet_scores(self, wb):
        ws = wb.create_sheet("4. Vendor Scores")
        hdr = ["SR#", "Vendor", "Facility", "Support", "Compliance", "Delivery", "Overall",
               "Blacklisted", "Blacklist Reason", "Notes for the AI evaluator", "Scored By"]
        ws.append(hdr)
        self._head(ws, hdr)
        qs = VendorScore.objects.select_related("vendor", "scored_by").order_by(
            "vendor__vendor_name")
        for i, s in enumerate(qs, start=1):
            sb = s.scored_by
            ws.append([i, s.vendor.vendor_name, s.facility_rating, s.support_rating,
                       s.compliance_rating, s.delivery_reliability, s.overall_rating,
                       yn(s.is_blacklisted), s.blacklist_reason, s.notes,
                       (sb.get_full_name() or sb.username) if sb else ""])
        self._finish(ws, 2, len(hdr), [6, 34, 11, 11, 12, 11, 10, 13, 46, 52, 18])
        for r in range(2, ws.max_row + 1):
            for c in range(3, 8):
                ws.cell(row=r, column=c).alignment = Alignment(
                    horizontal="center", vertical="center")
        self._pill(ws, hdr.index("Blacklisted") + 1, 2, {"Yes": RED, "No": GREEN})
        return ws.max_row - 1

    # ── 5. Requests ──────────────────────────────────────────────────────
    def sheet_requests(self, wb):
        ws = wb.create_sheet("5. Requests")
        hdr = ["Request ID", "Subject", "Country", "Product Verticals", "Scope of Work",
               "Warehouse Space", "Storage Temp", "Monthly Shipments", "Inbound", "Outbound",
               "Raised By", "Recipients", "Request Status", "Created At", "Bid Opens",
               "Bid Closes", "Bid Window", "Bids Received"]
        ws.append(hdr)
        self._head(ws, hdr)
        import re
        for vr in VendorRequest.objects.select_related("raised_by").order_by("id"):
            rb = vr.raised_by
            scope = re.sub(r"<[^>]+>", " ", vr.scope_of_work or "")
            scope = re.sub(r"\s+", " ", scope).strip()
            ws.append([
                vr.pk, vr.subject, vr.country, vr.product_verticals, scope,
                vr.warehouse_space_range, vr.storage_temp_range, vr.monthly_shipment_range,
                vr.inbound_range, vr.outbound_range,
                (rb.get_full_name() or rb.username) if rb else "",
                vr.recipient_count, vr.status,
                vr.created_at.strftime("%Y-%m-%d %H:%M"),
                vr.bid_open_at.strftime("%Y-%m-%d %H:%M") if vr.bid_open_at else "",
                vr.bid_close_at.strftime("%Y-%m-%d %H:%M") if vr.bid_close_at else "",
                vr.bid_window_status,
                VendorBid.objects.filter(vendor_request=vr).count(),
            ])
        self._finish(ws, 2, len(hdr),
                     [11, 46, 20, 30, 66, 16, 14, 18, 14, 14, 16, 12, 16, 18, 18, 18, 14, 14])
        self._pill(ws, hdr.index("Bid Window") + 1, 2,
                   {"open": GREEN, "closed": GREY, "not_started": AMBER})
        self._pill(ws, hdr.index("Request Status") + 1, 2,
                   {"sent": GREEN, "no_recipients": RED})
        return ws.max_row - 1

    # ── 6. Bids ──────────────────────────────────────────────────────────
    def sheet_bids(self, wb):
        ws = wb.create_sheet("6. Bids")
        hdr = ["Request ID", "Request", "Vendor", "Bid #", "Currency", "Storage", "Inbound",
               "Outbound", "Pick", "Packing", "Last-Mile", "Manpower", "TOTAL",
               "Staff Selected", "Status", "Notes / Message", "Submitted At"]
        ws.append(hdr)
        self._head(ws, hdr)
        qs = (VendorBid.objects
              .select_related("vendor_request", "vendor_registration", "submitted_by")
              .order_by("vendor_request_id", "vendor_registration__vendor_name", "bid_number"))
        for b in qs:
            reg = b.vendor_registration
            name = reg.vendor_name if reg else (
                b.submitted_by.get_full_name() or b.submitted_by.username)
            ws.append([b.vendor_request_id, b.vendor_request.subject, name, b.bid_number,
                       b.currency] +
                      [getattr(b, f) or Decimal(0) for f in COST_FIELDS] +
                      [bid_total(b), yn(b.is_selected), b.status, b.notes,
                       b.submitted_at.strftime("%Y-%m-%d %H:%M")])
        self._finish(ws, 2, len(hdr),
                     [11, 44, 34, 8, 10, 13, 12, 12, 11, 11, 12, 12, 15, 14, 13, 56, 18])
        first, total = hdr.index("Storage") + 1, hdr.index("TOTAL") + 1
        for r in range(2, ws.max_row + 1):
            for c in range(first, total + 1):
                cell = ws.cell(row=r, column=c)
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            t = ws.cell(row=r, column=total)
            t.font = Font(bold=True, size=11, color=NAVY)
            t.fill = PatternFill("solid", fgColor="EFF6FF")
            ws.cell(row=r, column=hdr.index("Bid #") + 1).alignment = Alignment(
                horizontal="center", vertical="center")
        self._pill(ws, hdr.index("Staff Selected") + 1, 2, {"Yes": GREEN, "No": "FFFFFF"},
                   default="FFFFFF")
        return ws.max_row - 1

    # ── 7. Data Dictionary ───────────────────────────────────────────────
    def sheet_dictionary(self, wb):
        ws = wb.create_sheet("7. Data Dictionary")
        hdr = ["Model", "Field", "Type", "Notes"]
        ws.append(hdr)
        self._head(ws, hdr)
        rows = [
            ("VendorRegistration", "vendor_name / country / state / city / warehouse_location", "text",
             "Submitted on the public form. state and city are 120 chars."),
            ("VendorRegistration", "email", "email",
             "OTP-verified before submit; duplicate pending/approved email refused."),
            ("VendorRegistration", "9 capability flags", "boolean", "Self-declared."),
            ("VendorRegistration", "status", "pending / approved / rejected", "Set by an admin."),
            ("Vendor", "identity + 9 capability flags", "mixed",
             "Created on approval — the Vendor Master record."),
            ("User (auth)", "username", "text", "Set to the vendor's email address."),
            ("UserProfile", "force_password_change", "boolean",
             "True until the vendor sets a new password; middleware enforces it."),
            ("VendorScore", "facility / support / compliance / delivery", "1-5",
             "Admin-set; overall_rating is their mean."),
            ("VendorScore", "is_blacklisted", "boolean",
             "Blacklisted vendors are ranked but can never be recommended."),
            ("VendorRequest", "scope_of_work", "HTML",
             "The evaluator reads the text to infer required services."),
            ("VendorRequest", "recipient_count / status", "int / sent|no_recipients",
             "Snapshot of who was emailed when the RFI went out."),
            ("VendorRequest", "bid_open_at / bid_close_at", "datetime",
             "Drives bid_window_status: open / not_started / closed."),
            ("VendorBid", "7 cost lines", "decimal(14,2)",
             "storage, inbound, outbound, pick, packing, last-mile, manpower. TOTAL is their sum."),
            ("VendorBid", "bid_number", "integer", "Revision number; vendors may re-quote."),
            ("VendorBid", "is_selected", "boolean",
             "Staff pick of the revision representing this vendor; the evaluator respects it."),
            ("VendorEvaluation", "(not exported)", "—",
             "Recommendation, reasoning, required services, decision points and the ranked "
             "list are produced live by OpenAI and shown in the app, not in this file."),
        ]
        for r in rows:
            ws.append(list(r))
        self._finish(ws, 2, len(hdr), [22, 48, 30, 76])
