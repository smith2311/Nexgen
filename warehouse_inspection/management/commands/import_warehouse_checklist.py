from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from warehouse_inspection.models import (
    WarehouseInspectionBuilding,
    ConditionCategory,
    WarehouseInspectionCheckpoint,
    WarehouseInspectionFloor,
    WarehouseInspectionFloorAssignment,
)


class Command(BaseCommand):
    help = 'Import warehouse condition/checklist/floor email masters from Check list.xlsx.'

    floor_columns = [
        ('G floor', 2, None),
        ('1st floor', 3, 7),
        ('2nd floor', 4, 8),
        ('3rd floor', 5, 9),
    ]

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            default='Check list.xlsx',
            help='Path to the checklist workbook. Defaults to Check list.xlsx in the project root.',
        )
        parser.add_argument(
            '--sheet',
            default='Sheet1',
            help='Worksheet name to import.',
        )
        parser.add_argument(
            '--created-by',
            default='scm.digitalsupport@merillife.com',
            help='Email to store as created_by for imported masters.',
        )
        parser.add_argument(
            '--building',
            default='Warehouse 1',
            help='Warehouse/building master name to import this checklist under.',
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError('openpyxl is required to import the warehouse checklist workbook.') from exc

        workbook_path = Path(options['file'])
        if not workbook_path.is_absolute():
            workbook_path = Path.cwd() / workbook_path
        if not workbook_path.exists():
            raise CommandError(f'Workbook not found: {workbook_path}')

        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        if options['sheet'] not in workbook.sheetnames:
            raise CommandError(f"Worksheet '{options['sheet']}' not found in {workbook_path.name}.")

        worksheet = workbook[options['sheet']]
        created_by = options['created_by']
        building_name = options['building']
        counters = {
            'buildings': 0,
            'floors': 0,
            'conditions': 0,
            'checkpoints': 0,
            'assignments': 0,
        }

        with transaction.atomic():
            building, building_created = WarehouseInspectionBuilding.objects.update_or_create(
                building_name=building_name,
                defaults={'building_order': 1, 'is_active': True},
            )
            if building_created:
                counters['buildings'] += 1

            floor_by_name = {}
            for order, (floor_name, _, _) in enumerate(self.floor_columns, start=1):
                floor, floor_created = WarehouseInspectionFloor.objects.update_or_create(
                    building=building,
                    floor_name=floor_name,
                    defaults={'floor_order': order, 'is_active': True},
                )
                floor_by_name[floor_name] = floor
                if floor_created:
                    counters['floors'] += 1

            current_condition = None
            checkpoint_order_by_condition = {}

            for row in worksheet.iter_rows(min_row=3, values_only=True):
                if not row or not any(cell is not None for cell in row):
                    continue

                condition_name = self._clean(row[0] if len(row) > 0 else None)
                checkpoint_text = self._clean(row[1] if len(row) > 1 else None)
                if condition_name:
                    current_condition = condition_name
                if not current_condition or not checkpoint_text:
                    continue

                condition, created = ConditionCategory.objects.update_or_create(
                    building=building,
                    condition_category_name=current_condition,
                    defaults={
                        'entity_type': 'admin',
                        'created_by': created_by,
                        'is_deleted': False,
                    },
                )
                if created:
                    counters['conditions'] += 1

                checkpoint_order_by_condition.setdefault(condition.condition_category_id, 0)
                checkpoint_order_by_condition[condition.condition_category_id] += 1
                checkpoint_order = checkpoint_order_by_condition[condition.condition_category_id]

                internal_email = self._clean(row[6] if len(row) > 6 else None)
                checkpoint, checkpoint_created = WarehouseInspectionCheckpoint.objects.update_or_create(
                    condition_category=condition,
                    checkpoint_name=checkpoint_text,
                    defaults={
                        'checkpoint_order': checkpoint_order,
                        'is_active': True,
                        'created_by': created_by,
                        'internal_department_email': internal_email or None,
                    },
                )
                if checkpoint_created:
                    counters['checkpoints'] += 1

                for floor_name, applicable_index, email_index in self.floor_columns:
                    applicable_value = self._clean(row[applicable_index] if len(row) > applicable_index else None)
                    if not applicable_value:
                        continue

                    responsible_email = ''
                    if email_index is not None:
                        responsible_email = self._clean(row[email_index] if len(row) > email_index else None)

                    _, assignment_created = WarehouseInspectionFloorAssignment.objects.update_or_create(
                        checkpoint=checkpoint,
                        floor=floor_by_name[floor_name],
                        defaults={
                            'is_applicable': True,
                            'responsible_email': responsible_email or None,
                        },
                    )
                    if assignment_created:
                        counters['assignments'] += 1

        self.stdout.write(
            self.style.SUCCESS(
                'Imported warehouse checklist masters: '
                f"{counters['buildings']} buildings, "
                f"{counters['floors']} floors, "
                f"{counters['conditions']} conditions, "
                f"{counters['checkpoints']} checkpoints, "
                f"{counters['assignments']} floor assignments."
            )
        )

    @staticmethod
    def _clean(value):
        if value is None:
            return ''
        return str(value).strip()
