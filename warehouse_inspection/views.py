"""
Warehouse Inspection Views
────────────────────────────────────────────────────────────────
REST API views for warehouse inspection operations.
"""
import json
import logging
from datetime import datetime

from rest_framework import generics, viewsets, status
from rest_framework.decorators import action
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.permissions import AllowAny, IsAuthenticated, SAFE_METHODS
from rest_framework.exceptions import PermissionDenied
from rest_framework.views import APIView
from django.conf import settings
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.db.models import Q
from django.db import IntegrityError, transaction
from django.db.models.deletion import ProtectedError
from django.utils import timezone
from django.http import JsonResponse
from django.views.decorators.http import require_http_methods
from django.core.exceptions import ValidationError
from django.core.validators import validate_email

from .models import (
    WarehouseInspectionBuilding,
    ConditionCategory,
    WarehouseInspectionCheckpoint,
    WarehouseInspectionFloor,
    WarehouseInspectionFloorAssignment,
    WarehouseInspectionLoginLog,
    WarehouseInspectionReportRecipient,
    WarehouseInspectionResponse,
    WarehouseInspectionSecondResponse,
    WarehouseInspectionSecondSubmission,
    WarehouseInspectionSubmission,
)
from .serializers import (
    WarehouseInspectionBuildingSerializer,
    ConditionCategorySerializer,
    ConditionCategoryDetailSerializer,
    ConditionCategoryCreateSerializer,
    WarehouseInspectionFloorSerializer,
    WarehouseInspectionFloorAssignmentSerializer,
    WarehouseInspectionCheckpointSerializer,
    WarehouseInspectionResponseSerializer,
    WarehouseInspectionResponsePhotoUploadSerializer,
    WarehouseInspectionSecondSubmissionSerializer,
    WarehouseInspectionSubmissionCreateSerializer,
    WarehouseInspectionSubmissionSerializer,
    get_warehouse_user_identity,
)
from .services import (
    send_inspection_submission_emails,
    send_inspection_summary_email,
    should_send_responsible_inspection_emails,
)


logger = logging.getLogger(__name__)

User = get_user_model()

WAREHOUSE_INSPECTION_ADMIN_GROUP = 'Warehouse Inspection Admin'
WAREHOUSE_INSPECTION_USER_GROUP = 'Warehouse Inspection User'
WAREHOUSE_INSPECTION_ROLE_GROUPS = (
    WAREHOUSE_INSPECTION_ADMIN_GROUP,
    WAREHOUSE_INSPECTION_USER_GROUP,
)
WAREHOUSE_INSPECTION_LEGACY_ADMIN_GROUPS = (
    'admin',
    'manager',
    'warehouse manager',
)
WAREHOUSE_INSPECTION_ALLOWED_EMAIL_DOMAIN = 'merillife.com'

# Cross-module super-admin group — members have full admin rights in both
# Warehouse Inspection and Material Rework without being Django superusers.
ROOT_USER_GROUP = 'Root User'
_WI_ERR_ACCESS_DENIED = 'Access denied.'
_WI_ACCESS_DENIED_TEMPLATE = 'warehouse_inspection/access_denied.html'


def _is_allowed_warehouse_email(email):
    return str(email or '').strip().lower().endswith(f'@{WAREHOUSE_INSPECTION_ALLOWED_EMAIL_DOMAIN}')


def is_warehouse_admin(user):
    """Return True for users allowed to maintain masters and email mappings."""
    if not user or not user.is_authenticated:
        return False
    if user.is_superuser or user.groups.filter(name=ROOT_USER_GROUP).exists():
        return True
    if not _is_allowed_warehouse_email(user.email):
        return False
    if user.is_staff:
        return True
    return user.groups.filter(
        Q(name__iexact=WAREHOUSE_INSPECTION_ADMIN_GROUP) |
        Q(name__iexact='admin') |
        Q(name__iexact='manager') |
        Q(name__iexact='warehouse manager')
    ).exists()


def is_warehouse_user(user):
    """Return True for users allowed to open and submit warehouse inspections."""
    if not user or not user.is_authenticated:
        return False
    if is_warehouse_admin(user):
        return True
    if not _is_allowed_warehouse_email(user.email):
        return False
    return user.groups.filter(name__iexact=WAREHOUSE_INSPECTION_USER_GROUP).exists()


def _ensure_warehouse_inspection_groups():
    """Create Warehouse Inspection role groups if they are missing."""
    groups = {}
    for group_name in WAREHOUSE_INSPECTION_ROLE_GROUPS:
        groups[group_name], _ = Group.objects.get_or_create(name=group_name)
    return groups


def _has_group_iexact(user, group_name):
    return user.groups.filter(name__iexact=group_name).exists()


def _has_legacy_admin_access(user):
    if user.is_staff:
        return True
    group_query = Q()
    for group_name in WAREHOUSE_INSPECTION_LEGACY_ADMIN_GROUPS:
        group_query |= Q(name__iexact=group_name)
    return user.groups.filter(group_query).exists()


def _report_recipient_for_email(email):
    if not email:
        return None
    return WarehouseInspectionReportRecipient.objects.filter(email__iexact=email).first()


def _serialize_warehouse_access_user(user, current_user, report_recipient_by_email=None):
    has_admin_role = _has_group_iexact(user, WAREHOUSE_INSPECTION_ADMIN_GROUP)
    has_user_role = _has_group_iexact(user, WAREHOUSE_INSPECTION_USER_GROUP)
    inherited_admin = _has_legacy_admin_access(user) and not has_admin_role
    is_allowed_domain = _is_allowed_warehouse_email(user.email)
    is_root = user.groups.filter(name=ROOT_USER_GROUP).exists()
    recipient = None
    recipient_key = str(user.email or '').strip().lower()
    if report_recipient_by_email is not None:
        recipient = report_recipient_by_email.get(recipient_key)
    else:
        recipient = _report_recipient_for_email(user.email)
    return {
        'id': user.id,
        'username': user.username,
        'full_name': user.get_full_name() or user.username,
        'email': user.email,
        'is_wi_admin': is_warehouse_admin(user),
        'is_wi_user': has_user_role and not is_warehouse_admin(user),
        'has_wi_admin_role': has_admin_role,
        'has_wi_user_role': has_user_role,
        'is_inherited_admin': inherited_admin,
        'is_root_user': is_root,
        'is_current_user': user.id == current_user.id,
        'is_staff': user.is_staff,
        'is_superuser': user.is_superuser,
        'is_allowed_domain': is_allowed_domain,
        'receives_report_emails': bool(recipient and recipient.is_active),
        'report_recipient_id': recipient.recipient_id if recipient else None,
    }


def _parse_json_body(request):
    try:
        return json.loads(request.body or b'{}')
    except (TypeError, json.JSONDecodeError):
        raise ValueError('Invalid request body.')


def _get_target_user(user_id):
    if not user_id:
        return None
    try:
        target_id = int(user_id)
    except (TypeError, ValueError):
        return None
    try:
        return User.objects.get(id=target_id)
    except User.DoesNotExist:
        return None


def _can_manage_warehouse_role(target_user, acting_user):
    if target_user.id == acting_user.id:
        return 'You cannot change your own Warehouse Inspection role here.'
    if target_user.is_superuser:
        return 'Superusers cannot be changed here.'
    # Root Users can only be managed by other Root Users or Django superusers
    if target_user.groups.filter(name=ROOT_USER_GROUP).exists():
        if not (acting_user.is_superuser or acting_user.groups.filter(name=ROOT_USER_GROUP).exists()):
            return 'Root Users can only be managed by other Root Users or Superusers.'
    if _has_legacy_admin_access(target_user) and not _has_group_iexact(target_user, WAREHOUSE_INSPECTION_ADMIN_GROUP):
        return 'This user has Warehouse Inspection admin access from a global staff/admin role.'
    return ''


def _domain_error():
    return f'Only @{WAREHOUSE_INSPECTION_ALLOWED_EMAIL_DOMAIN} email addresses can access Warehouse Inspection.'


def _apply_warehouse_role(user, role):
    role_groups = _ensure_warehouse_inspection_groups()
    admin_group = role_groups[WAREHOUSE_INSPECTION_ADMIN_GROUP]
    user_group = role_groups[WAREHOUSE_INSPECTION_USER_GROUP]
    user.groups.remove(admin_group, user_group)
    if role == 'admin':
        user.groups.add(admin_group)
    elif role == 'user':
        user.groups.add(user_group)


def _unique_username_from_email(email):
    base_username = (email.split('@', 1)[0] or 'warehouse_user')[:140]
    username = base_username
    counter = 1
    while User.objects.filter(username=username).exists():
        suffix = str(counter)
        username = f'{base_username[:150 - len(suffix)]}{suffix}'
        counter += 1
    return username


class WarehouseReadAdminWriteMixin:
    """Allow inspection users to read masters, but only admins can change them."""

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not is_warehouse_user(request.user):
            raise PermissionDenied('You do not have access to Warehouse Inspection.')
        if request.method not in SAFE_METHODS and not is_warehouse_admin(request.user):
            raise PermissionDenied('Only admins can modify warehouse inspection setup.')


class WarehouseSubmissionAccessMixin:
    """Allow inspection users to submit inspections and read their own history."""

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not is_warehouse_user(request.user):
            raise PermissionDenied('You do not have access to Warehouse Inspection.')
        if (
            request.method not in SAFE_METHODS
            and request.method != 'POST'
            and not is_warehouse_admin(request.user)
        ):
            raise PermissionDenied('Only admins can modify submitted inspections.')


def mark_submission_email_error(submissions, exc):
    """Persist unexpected notification errors without failing saved submissions."""
    submissions = list(submissions)
    message = f'Inspection saved, but email notification failed: {exc}'
    for submission in submissions:
        submission.email_sent = False
        submission.email_error = message
        submission.save(update_fields=['email_sent', 'email_error'])
    return {
        'recipients': [],
        'cc': [],
        'sent_count': 0,
        'failed_count': len(submissions),
        'errors': [message],
    }


def clean_unique_values(values):
    """Return stable unique non-empty values for API summaries."""
    seen = set()
    cleaned = []
    for value in values:
        value = str(value or '').strip()
        if not value:
            continue
        key = value.lower()
        if key in seen:
            continue
        seen.add(key)
        cleaned.append(value)
    return cleaned


def merge_email_results(results):
    """Combine notification summaries from first and second inspection tables."""
    results = [result for result in results if result]
    return {
        'recipients': clean_unique_values(
            email
            for result in results
            for email in result.get('recipients', [])
        ),
        'cc': clean_unique_values(
            email
            for result in results
            for email in result.get('cc', [])
        ),
        'sent_count': sum(int(result.get('sent_count', 0) or 0) for result in results),
        'failed_count': sum(int(result.get('failed_count', 0) or 0) for result in results),
        'errors': [
            error
            for result in results
            for error in result.get('errors', [])
        ],
    }


def send_notifications_for_submissions(submissions):
    """Send configured warehouse inspection notifications."""
    submissions = list(submissions)
    primary_submissions = [
        submission for submission in submissions
        if isinstance(submission, WarehouseInspectionSubmission)
    ]
    second_submissions = [
        submission for submission in submissions
        if isinstance(submission, WarehouseInspectionSecondSubmission)
    ]
    results = []
    if should_send_responsible_inspection_emails():
        if primary_submissions:
            results.append(send_inspection_submission_emails(primary_submissions))
        if second_submissions:
            results.append(send_inspection_submission_emails(second_submissions))
    results.append(send_inspection_summary_email(submissions))
    return merge_email_results(results)


def serialize_submission_instance(submission, context):
    """Serialize a first or second inspection submission for API responses."""
    serializer_class = (
        WarehouseInspectionSecondSubmissionSerializer
        if isinstance(submission, WarehouseInspectionSecondSubmission)
        else WarehouseInspectionSubmissionSerializer
    )
    return serializer_class(submission, context=context).data


def submission_integrity_error_response(exc):
    """Return a friendly response for duplicate/constraint errors during submit."""
    message = str(exc)
    if 'wh_submission_daily_unique' in message:
        detail = (
            "This floor's inspection for today is being saved by another request. "
            'Please retry.'
        )
    else:
        detail = 'Inspection could not be saved because of a database constraint.'
    return Response(
        {
            'detail': detail,
            'error': message,
        },
        status=status.HTTP_409_CONFLICT,
    )


class WarehouseInspectionBuildingViewSet(WarehouseReadAdminWriteMixin, viewsets.ModelViewSet):
    """ViewSet for warehouse/building masters."""

    permission_classes = [IsAuthenticated]
    serializer_class = WarehouseInspectionBuildingSerializer

    def get_queryset(self):
        queryset = WarehouseInspectionBuilding.objects.all()
        if not is_warehouse_admin(self.request.user):
            queryset = queryset.filter(is_active=True)
        return queryset

    def destroy(self, request, *args, **kwargs):
        """Delete a building and its warehouse inspection setup/history."""
        building = self.get_object()
        try:
            self._delete_building_completely(building)
        except ProtectedError:
            return Response(
                {
                    'detail': (
                        'Could not delete this building because related records are still protected.'
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )
        return Response(status=status.HTTP_204_NO_CONTENT)

    def _delete_building_completely(self, building):
        """Remove a building with all dependent warehouse-inspection data."""
        with transaction.atomic():
            WarehouseInspectionResponse.objects.filter(
                Q(submission__building=building)
                | Q(floor__building=building)
                | Q(checkpoint__condition_category__building=building)
            ).delete()
            WarehouseInspectionSecondResponse.objects.filter(
                Q(submission__building=building)
                | Q(floor__building=building)
                | Q(checkpoint__condition_category__building=building)
            ).delete()
            WarehouseInspectionSecondSubmission.objects.filter(building=building).delete()
            WarehouseInspectionSubmission.objects.filter(building=building).delete()
            ConditionCategory.objects.filter(building=building).delete()
            WarehouseInspectionFloor.objects.filter(building=building).delete()
            building.delete()


class ConditionCategoryViewSet(WarehouseReadAdminWriteMixin, viewsets.ModelViewSet):
    """
    API ViewSet for managing Condition Categories.
    Supports CRUD operations with role-based filtering.
    """
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        """All authenticated users can select active condition masters."""
        queryset = ConditionCategory.objects.select_related('building').filter(is_deleted=False)
        building_id = self.request.query_params.get('building_id')
        if building_id:
            queryset = queryset.filter(building_id=building_id)
        if not is_warehouse_admin(self.request.user):
            queryset = queryset.filter(building__is_active=True)
        return queryset

    def get_serializer_class(self):
        """Use different serializers based on action."""
        if self.action == 'retrieve':
            return ConditionCategoryDetailSerializer
        elif self.action == 'create':
            return ConditionCategoryCreateSerializer
        else:
            return ConditionCategorySerializer

    def create(self, request, *args, **kwargs):
        """Create a condition category."""
        building_id = request.data.get('building')
        condition_name = str(request.data.get('condition_category_name', '')).strip()
        if building_id and condition_name:
            existing_condition = (
                ConditionCategory.objects
                .filter(
                    building_id=building_id,
                    condition_category_name__iexact=condition_name,
                    is_deleted=False,
                )
                .first()
            )
            if existing_condition:
                serializer = ConditionCategorySerializer(
                    existing_condition,
                    context=self.get_serializer_context(),
                )
                data = serializer.data
                data['already_exists'] = True
                data['detail'] = 'Condition category already exists.'
                return Response(data, status=status.HTTP_200_OK)

        return super().create(request, *args, **kwargs)

    def perform_create(self, serializer):
        """Save with logged-in user email."""
        email, _ = get_warehouse_user_identity(self.request.user)
        serializer.save(created_by=email)

    def perform_update(self, serializer):
        """Update condition category."""
        serializer.save()

    def destroy(self, request, *args, **kwargs):
        """
        Hard delete a condition category from DB.

        Notes:
        - Category delete should remove it completely from the database.
        - Existing inspection submissions are deleted first because the
          submission model protects its condition/floor masters.
        - Individual checkpoints/items should be "hidden" using is_active=False
          (handled in the checkpoint viewset / UI save flow).
        """
        condition = self.get_object()
        self._delete_condition_completely(condition)
        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(detail=True, methods=['post'])
    def soft_delete(self, request, pk=None):
        """UI-friendly endpoint that performs a hard delete."""
        condition = self.get_object()
        self._delete_condition_completely(condition)
        return Response({'detail': 'Condition category deleted.'}, status=status.HTTP_200_OK)

    def _delete_condition_completely(self, condition):
        """Remove the category and its dependent inspection data."""
        with transaction.atomic():
            WarehouseInspectionSecondSubmission.objects.filter(
                condition_category=condition
            ).delete()
            WarehouseInspectionSubmission.objects.filter(
                condition_category=condition
            ).delete()
            condition.delete()

    @action(detail=True, methods=['post'])
    def restore(self, request, pk=None):
        """
        Restore is not supported when categories are hard-deleted.
        Kept for compatibility with older admin UI actions.
        """
        return Response(
            {'detail': 'Restore is not available because categories are permanently deleted.'},
            status=status.HTTP_400_BAD_REQUEST,
        )


class WarehouseInspectionFloorViewSet(WarehouseReadAdminWriteMixin, viewsets.ModelViewSet):
    """ViewSet for warehouse floor masters."""

    permission_classes = [IsAuthenticated]
    serializer_class = WarehouseInspectionFloorSerializer

    def get_queryset(self):
        queryset = WarehouseInspectionFloor.objects.select_related('building').all()
        building_id = self.request.query_params.get('building_id')
        if building_id:
            queryset = queryset.filter(building_id=building_id)
        if not is_warehouse_admin(self.request.user):
            queryset = queryset.filter(is_active=True, building__is_active=True)
        return queryset


class WarehouseInspectionFloorAssignmentViewSet(WarehouseReadAdminWriteMixin, viewsets.ModelViewSet):
    """ViewSet for floor/checkpoint responsible-person mappings."""

    permission_classes = [IsAuthenticated]
    serializer_class = WarehouseInspectionFloorAssignmentSerializer

    def get_queryset(self):
        queryset = WarehouseInspectionFloorAssignment.objects.select_related(
            'checkpoint',
            'checkpoint__condition_category',
            'checkpoint__condition_category__building',
            'floor',
            'floor__building',
        )
        checkpoint_id = self.request.query_params.get('checkpoint_id')
        floor_id = self.request.query_params.get('floor_id')
        condition_id = self.request.query_params.get('condition_id')
        building_id = self.request.query_params.get('building_id')
        if checkpoint_id:
            queryset = queryset.filter(checkpoint_id=checkpoint_id)
        if floor_id:
            queryset = queryset.filter(floor_id=floor_id)
        if condition_id:
            queryset = queryset.filter(checkpoint__condition_category_id=condition_id)
        if building_id:
            queryset = queryset.filter(floor__building_id=building_id)
        if not is_warehouse_admin(self.request.user):
            queryset = queryset.filter(
                is_applicable=True,
                floor__is_active=True,
                floor__building__is_active=True,
            )
        return queryset


class WarehouseInspectionCheckpointViewSet(WarehouseReadAdminWriteMixin, viewsets.ModelViewSet):
    """ViewSet for managing inspection checkpoints."""

    permission_classes = [IsAuthenticated]
    serializer_class = WarehouseInspectionCheckpointSerializer

    def get_queryset(self):
        """Filter checkpoints by condition category."""
        queryset = WarehouseInspectionCheckpoint.objects.select_related(
            'condition_category',
            'condition_category__building',
        ).prefetch_related(
            'floor_assignments',
            'floor_assignments__floor',
        ).filter(is_active=True)
        building_id = self.request.query_params.get('building_id')
        condition_id = self.request.query_params.get('condition_id')
        floor_id = self.request.query_params.get('floor_id')
        if building_id:
            queryset = queryset.filter(condition_category__building_id=building_id)
        if condition_id:
            queryset = queryset.filter(condition_category_id=condition_id)
        if floor_id:
            queryset = queryset.filter(
                floor_assignments__floor_id=floor_id,
                floor_assignments__is_applicable=True,
            ).distinct()
        return queryset

    def perform_create(self, serializer):
        """Save with logged-in user email."""
        email, _ = get_warehouse_user_identity(self.request.user)
        serializer.save(created_by=email)

    def destroy(self, request, *args, **kwargs):
        """Hide checklist items instead of permanently deleting them."""
        checkpoint = self.get_object()
        checkpoint.is_active = False
        checkpoint.save(update_fields=['is_active', 'updated_at'])
        return Response(status=status.HTTP_204_NO_CONTENT)


class WarehouseInspectionResponseViewSet(WarehouseReadAdminWriteMixin, viewsets.ReadOnlyModelViewSet):
    """Read-only response history. Users submit responses through submissions."""

    permission_classes = [IsAuthenticated]
    serializer_class = WarehouseInspectionResponseSerializer

    def get_queryset(self):
        """Get responses for the logged-in user."""
        queryset = WarehouseInspectionResponse.objects.select_related(
            'submission', 'checkpoint', 'floor', 'floor_assignment'
        )
        if is_warehouse_admin(self.request.user):
            return queryset
        return queryset.filter(responded_by_email=self.request.user.email)

    @action(detail=False, methods=['get'])
    def by_checkpoint(self, request):
        """Get all responses for a specific checkpoint."""
        checkpoint_id = request.query_params.get('checkpoint_id')
        if not checkpoint_id:
            return Response(
                {'error': 'checkpoint_id required'},
                status=status.HTTP_400_BAD_REQUEST
            )

        responses = WarehouseInspectionResponse.objects.filter(
            checkpoint_id=checkpoint_id
        )
        if not is_warehouse_admin(request.user):
            responses = responses.filter(responded_by_email=request.user.email)
        serializer = self.get_serializer(responses, many=True)
        return Response(serializer.data)


class WarehouseInspectionResponsePhotoUploadView(generics.CreateAPIView):
    """Accepts one checkpoint proof photo, uploaded ahead of the final checklist submit."""

    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    serializer_class = WarehouseInspectionResponsePhotoUploadSerializer

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not is_warehouse_user(request.user):
            raise PermissionDenied('You do not have access to Warehouse Inspection.')

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        upload = serializer.save()
        return Response(
            {
                'upload_token': upload.upload_token,
                'url': request.build_absolute_uri(upload.photo.url),
            },
            status=status.HTTP_201_CREATED,
        )


class WarehouseInspectionSubmissionViewSet(WarehouseSubmissionAccessMixin, viewsets.ModelViewSet):
    """ViewSet for completed warehouse inspection submissions."""

    permission_classes = [IsAuthenticated]

    def get_serializer_class(self):
        if self.action == 'create':
            return WarehouseInspectionSubmissionCreateSerializer
        return WarehouseInspectionSubmissionSerializer

    def get_queryset(self):
        queryset = WarehouseInspectionSubmission.objects.select_related(
            'building',
            'condition_category',
            'floor',
        ).prefetch_related(
            'responses',
            'responses__checkpoint',
            'responses__floor',
        )
        building_id = self.request.query_params.get('building_id')
        if building_id:
            queryset = queryset.filter(building_id=building_id)
        if is_warehouse_admin(self.request.user):
            return queryset
        return queryset.filter(submitted_by_email=self.request.user.email)

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            submission = serializer.save()
        except IntegrityError as exc:
            return submission_integrity_error_response(exc)
        try:
            send_notifications_for_submissions([submission])
        except Exception as exc:
            mark_submission_email_error([submission], exc)
            submission.refresh_from_db(
                fields=['email_sent', 'email_sent_at', 'email_error']
            )
        data = serialize_submission_instance(submission, self.get_serializer_context())
        headers = self.get_success_headers(data)
        return Response(data, status=status.HTTP_201_CREATED, headers=headers)

    @action(detail=False, methods=['post'], url_path='batch')
    def batch(self, request):
        items = request.data.get('submissions')
        if not isinstance(items, list) or not items:
            return Response(
                {'submissions': 'At least one floor submission is required.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        floor_ids = [str(item.get('floor')) for item in items if item.get('floor')]
        if len(floor_ids) != len(set(floor_ids)):
            return Response(
                {'submissions': 'Duplicate floor submissions are not allowed.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        serializers = []
        errors = {}
        for index, item in enumerate(items):
            data = {
                'building': item.get('building') or request.data.get('building'),
                'condition_category': (
                    item.get('condition_category')
                    or request.data.get('condition_category')
                ),
                'floor': item.get('floor'),
                'notes': item.get('notes', ''),
                'responses': item.get('responses', []),
            }
            serializer = WarehouseInspectionSubmissionCreateSerializer(
                data=data,
                context=self.get_serializer_context(),
            )
            if serializer.is_valid():
                serializers.append(serializer)
            else:
                errors[index] = serializer.errors

        if errors:
            return Response(
                {'submissions': errors},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            with transaction.atomic():
                submissions = [serializer.save() for serializer in serializers]
        except IntegrityError as exc:
            return submission_integrity_error_response(exc)

        try:
            email_result = send_notifications_for_submissions(submissions)
        except Exception as exc:
            email_result = mark_submission_email_error(submissions, exc)

        data = [
            serialize_submission_instance(submission, self.get_serializer_context())
            for submission in submissions
        ]
        headers = self.get_success_headers(data)
        return Response(
            {
                'submissions': data,
                'email': email_result,
            },
            status=status.HTTP_201_CREATED,
            headers=headers,
        )


class WarehouseInspectionSecondSubmissionViewSet(
    WarehouseSubmissionAccessMixin,
    viewsets.ReadOnlyModelViewSet,
):
    """Read-only API for rare second same-day inspection submissions."""

    permission_classes = [IsAuthenticated]
    serializer_class = WarehouseInspectionSecondSubmissionSerializer

    def get_queryset(self):
        queryset = WarehouseInspectionSecondSubmission.objects.select_related(
            'first_submission',
            'building',
            'condition_category',
            'floor',
        ).prefetch_related(
            'responses',
            'responses__checkpoint',
            'responses__floor',
        )
        building_id = self.request.query_params.get('building_id')
        if building_id:
            queryset = queryset.filter(building_id=building_id)
        if is_warehouse_admin(self.request.user):
            return queryset
        return queryset.filter(submitted_by_email=self.request.user.email)


# ────────────────────────────────────────────────────────────────
# Warehouse Inspection Admin Access API
# ────────────────────────────────────────────────────────────────

@login_required
@require_http_methods(['GET'])
def admin_user_list(request):
    """Return active users and their Warehouse Inspection role state."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    role_groups = _ensure_warehouse_inspection_groups()
    access_query = (
        Q(groups__in=[role_groups[WAREHOUSE_INSPECTION_ADMIN_GROUP], role_groups[WAREHOUSE_INSPECTION_USER_GROUP]]) |
        Q(is_staff=True) |
        Q(groups__name=ROOT_USER_GROUP)
    )
    for group_name in WAREHOUSE_INSPECTION_LEGACY_ADMIN_GROUPS:
        access_query |= Q(groups__name__iexact=group_name)

    users = (
        User.objects
        .prefetch_related('groups')
        .filter(is_active=True, is_superuser=False)
        .filter(access_query)
        .distinct()
        .order_by('first_name', 'last_name', 'username')
    )
    recipient_by_email = {
        recipient.email.lower(): recipient
        for recipient in WarehouseInspectionReportRecipient.objects.all()
    }
    return JsonResponse({
        'users': [
            _serialize_warehouse_access_user(user, request.user, recipient_by_email)
            for user in users
        ],
    })


@login_required
@require_http_methods(['POST'])
def update_user_role(request):
    """Set a user's Warehouse Inspection role to admin, user, or no access."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    try:
        data = _parse_json_body(request)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)

    target_user = _get_target_user(data.get('user_id'))
    if target_user is None:
        return JsonResponse({'error': 'User not found.'}, status=404)

    management_error = _can_manage_warehouse_role(target_user, request.user)
    if management_error:
        return JsonResponse({'error': management_error}, status=400)

    role = str(data.get('role', '')).strip().lower()
    if role not in {'admin', 'user', 'none'}:
        return JsonResponse({'error': 'role must be admin, user, or none.'}, status=400)
    if role in {'admin', 'user'} and not _is_allowed_warehouse_email(target_user.email):
        return JsonResponse({'error': _domain_error()}, status=400)

    with transaction.atomic():
        _apply_warehouse_role(target_user, role)
        if role != 'admin':
            WarehouseInspectionReportRecipient.objects.filter(
                email__iexact=target_user.email
            ).delete()

    target_user = User.objects.prefetch_related('groups').get(id=target_user.id)
    label = target_user.get_full_name() or target_user.username
    return JsonResponse({
        'success': True,
        'removed': role == 'none',
        'user': _serialize_warehouse_access_user(target_user, request.user),
        'message': f'Updated Warehouse Inspection role for {label}.',
    })


def _resolve_warehouse_user(email, existing_user, role):
    """Get-or-create a user for WI access. Returns (user, created, restored, error_str)."""
    if existing_user:
        target = existing_user
        if target.is_superuser:
            return None, False, False, 'Superusers cannot be changed here.'
        if _has_legacy_admin_access(target) and not _has_group_iexact(target, WAREHOUSE_INSPECTION_ADMIN_GROUP):
            return None, False, False, (
                'This user already has Warehouse Inspection admin access from a global staff/admin role.'
            )
        restored = False
        if not target.is_active:
            target.is_active = True
            target.save(update_fields=['is_active'])
            restored = True
        _apply_warehouse_role(target, role)
        return target, False, restored, None

    target = User.objects.create_user(
        username=_unique_username_from_email(email),
        email=email,
        password=None,
    )
    _apply_warehouse_role(target, role)
    return target, True, False, None


@login_required
@require_http_methods(['POST'])
def create_user_access(request):
    """Create or restore a user by email, then assign Warehouse Inspection access."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    try:
        data = _parse_json_body(request)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)

    email = User.objects.normalize_email(str(data.get('email', '')).strip())
    role = str(data.get('role', 'user')).strip().lower()
    if role not in {'admin', 'user'}:
        return JsonResponse({'error': 'role must be admin or user.'}, status=400)
    if not email:
        return JsonResponse({'error': 'Email is required.'}, status=400)
    try:
        validate_email(email)
    except ValidationError:
        return JsonResponse({'error': 'Enter a valid email address.'}, status=400)
    if not _is_allowed_warehouse_email(email):
        return JsonResponse({'error': _domain_error()}, status=400)

    existing_user = User.objects.filter(email__iexact=email).first()
    with transaction.atomic():
        target_user, created, restored, error = _resolve_warehouse_user(email, existing_user, role)
    if error:
        return JsonResponse({'error': error}, status=400)

    target_user = User.objects.prefetch_related('groups').get(id=target_user.id)
    label = target_user.get_full_name() or target_user.username
    if created:
        message = f'Added {label} to Warehouse Inspection.'
    elif restored:
        message = f'Restored {label} and updated Warehouse Inspection access.'
    else:
        message = f'Updated Warehouse Inspection access for {label}.'

    return JsonResponse({
        'success': True,
        'created': created,
        'restored': restored,
        'user': _serialize_warehouse_access_user(target_user, request.user),
        'message': message,
    })


@login_required
@require_http_methods(['POST'])
def remove_user_access(request):
    """Remove Warehouse Inspection access without deleting the account."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    try:
        data = _parse_json_body(request)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)

    target_user = _get_target_user(data.get('user_id'))
    if target_user is None:
        return JsonResponse({'error': 'User not found.'}, status=404)

    management_error = _can_manage_warehouse_role(target_user, request.user)
    if management_error:
        return JsonResponse({'error': management_error}, status=400)

    role_groups = _ensure_warehouse_inspection_groups()
    admin_group = role_groups[WAREHOUSE_INSPECTION_ADMIN_GROUP]
    user_group = role_groups[WAREHOUSE_INSPECTION_USER_GROUP]
    label = target_user.get_full_name() or target_user.username
    with transaction.atomic():
        target_user.groups.remove(admin_group, user_group)
        WarehouseInspectionReportRecipient.objects.filter(
            email__iexact=target_user.email
        ).delete()

    target_user = User.objects.prefetch_related('groups').get(id=target_user.id)
    return JsonResponse({
        'success': True,
        'removed': True,
        'user': _serialize_warehouse_access_user(target_user, request.user),
        'message': f'Removed Warehouse Inspection access for {label}.',
    })


# ────────────────────────────────────────────────────────────────
# Report Recipients Master API
# (daily activity report + inspection summary email recipients)
# ────────────────────────────────────────────────────────────────

def _serialize_report_recipient(recipient):
    return {
        'id': recipient.recipient_id,
        'email': recipient.email,
        'name': recipient.name,
        'is_active': recipient.is_active,
        'created_by': recipient.created_by,
        'created_at': recipient.created_at.isoformat(),
    }


@login_required
@require_http_methods(['GET'])
def report_recipients_list(request):
    """Return everyone in the Report Recipients master."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    recipients = WarehouseInspectionReportRecipient.objects.all()
    return JsonResponse({
        'recipients': [_serialize_report_recipient(recipient) for recipient in recipients],
    })


@login_required
@require_http_methods(['POST'])
def add_report_recipient(request):
    """Add (or reactivate) an email in the Report Recipients master."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    try:
        data = _parse_json_body(request)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)

    email = User.objects.normalize_email(str(data.get('email', '')).strip())
    name = str(data.get('name', '')).strip()
    if not email:
        return JsonResponse({'error': 'Email is required.'}, status=400)
    try:
        validate_email(email)
    except ValidationError:
        return JsonResponse({'error': 'Enter a valid email address.'}, status=400)

    user = User.objects.filter(email__iexact=email, is_active=True).first()
    if not user or not is_warehouse_admin(user):
        return JsonResponse(
            {'error': 'Only active Warehouse Inspection admins can receive report emails.'},
            status=400,
        )

    recipient = WarehouseInspectionReportRecipient.objects.filter(email__iexact=email).first()
    if recipient:
        created = False
        update_fields = []
        if not recipient.is_active:
            recipient.is_active = True
            update_fields.append('is_active')
        if name and recipient.name != name:
            recipient.name = name
            update_fields.append('name')
        if update_fields:
            recipient.save(update_fields=update_fields)
        message = f'{recipient.email} is already in Report Recipients.'
    else:
        recipient = WarehouseInspectionReportRecipient.objects.create(
            email=email,
            name=name,
            created_by=request.user.email,
        )
        created = True
        message = f'Added {recipient.email} to Report Recipients.'

    return JsonResponse({
        'success': True,
        'created': created,
        'recipient': _serialize_report_recipient(recipient),
        'message': message,
    })


@login_required
@require_http_methods(['POST'])
def remove_report_recipient(request):
    """Remove an email from the Report Recipients master."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    try:
        data = _parse_json_body(request)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)

    recipients = WarehouseInspectionReportRecipient.objects.all()
    if data.get('id'):
        recipients = recipients.filter(recipient_id=data.get('id'))
    else:
        email = User.objects.normalize_email(str(data.get('email', '')).strip())
        recipients = recipients.filter(email__iexact=email)

    deleted_count, _ = recipients.delete()
    if not deleted_count:
        return JsonResponse({'error': 'Recipient not found.'}, status=404)

    return JsonResponse({'success': True, 'removed': True})


# ────────────────────────────────────────────────────────────────
# Frontend Views (HTML)
# ────────────────────────────────────────────────────────────────

WI_LOGIN_SESSION_KEY = 'wi_login_logged_date'


def record_warehouse_login(request):
    """
    Log one WarehouseInspectionLoginLog row the first time a given user opens
    the module on a given day (session-flag-guarded so it's once-per-day per
    user, not once-per-request). Feeds the "who logged in and when" section of
    the daily activity report - see services.build_daily_warehouse_inspection_report.
    """
    today_iso = timezone.localdate().isoformat()
    if request.session.get(WI_LOGIN_SESSION_KEY) == today_iso:
        return
    email, name = get_warehouse_user_identity(request.user)
    WarehouseInspectionLoginLog.objects.create(user_email=email, user_name=name)
    request.session[WI_LOGIN_SESSION_KEY] = today_iso


@login_required
def condition_management(request):
    """Admin condition management page."""
    user = request.user
    is_admin = is_warehouse_admin(user)

    if not is_admin:
        return render(request, _WI_ACCESS_DENIED_TEMPLATE)

    context = {
        'page_title': 'Condition Categories',
        'user': user,
        'is_admin': is_admin,
    }
    return render(request, 'warehouse_inspection/admin_conditions.html', context)


@login_required
def report_recipients_page(request):
    """Admin page for managing who receives report emails (daily report + summary)."""
    user = request.user
    is_admin = is_warehouse_admin(user)

    if not is_admin:
        return render(request, _WI_ACCESS_DENIED_TEMPLATE)

    context = {
        'page_title': 'Report Recipients',
        'user': user,
        'is_admin': is_admin,
    }
    return render(request, 'warehouse_inspection/report_recipients.html', context)


@login_required
def inspection_dashboard(request):
    """
    Warehouse inspection submodule selector.

    Admins see the full module-selector dashboard.
    Regular (non-admin) users are redirected straight to the inspection form
    because the dashboard would only show them a single card anyway.
    """
    from django.shortcuts import redirect as _redirect

    user = request.user
    is_admin = is_warehouse_admin(user)
    can_inspect = is_warehouse_user(user)

    if not can_inspect:
        return render(request, _WI_ACCESS_DENIED_TEMPLATE)

    record_warehouse_login(request)

    # Non-admin WI users have no reason to land on the module-selector;
    # send them straight to the inspection form.
    if not is_admin:
        return _redirect('warehouse_inspection:inspection_form')

    context = {
        'page_title': 'Warehouse Inspection',
        'user': user,
        'is_admin': is_admin,
        'can_inspect': can_inspect,
    }
    return render(request, 'warehouse_inspection/dashboard.html', context)


@login_required
def inspection_form(request):
    """Condition/floor inspection form."""
    user = request.user
    is_admin = is_warehouse_admin(user)
    can_inspect = is_warehouse_user(user)
    if not can_inspect:
        return render(request, _WI_ACCESS_DENIED_TEMPLATE)

    record_warehouse_login(request)

    context = {
        'page_title': 'Warehouse Inspection Checklist',
        'user': user,
        'is_admin': is_admin,
        'can_inspect': can_inspect,
    }
    return render(request, 'warehouse_inspection/inspection_form.html', context)


# ────────────────────────────────────────────────────────────────
# Excel Checklist Import
# ────────────────────────────────────────────────────────────────

def _find_email_separator_col(headers):
    for i, h in enumerate(headers):
        if 'internal' in h.lower() or 'email id' in h.lower():
            return i
    return 6


def _row1_building_by_col(worksheet, first_col, last_col):
    """
    Map each column (0-indexed, matching the `headers`/row-tuple convention
    used elsewhere in this module) in [first_col, last_col) to the building
    name from row 1 that covers it.

    Real checklist workbooks group several floor columns for one building
    under a single merged row-1 cell (e.g. "Meril Park 3" merged across 4
    floor columns), while a building with only one floor gets an unmerged
    single-cell name (e.g. "Meril Park" over just "Ground Floor"). A blank,
    unmerged cell falls back to the nearest building name to its left, so a
    sheet that repeats/omits the name for adjacent same-building columns
    still resolves correctly.
    """
    building_by_col = {}
    for merged_range in worksheet.merged_cells.ranges:
        if merged_range.min_row != 1 or merged_range.max_row != 1:
            continue
        value = str(worksheet.cell(row=1, column=merged_range.min_col).value or '').strip()
        if not value:
            continue
        for col_1indexed in range(merged_range.min_col, merged_range.max_col + 1):
            building_by_col[col_1indexed - 1] = value

    row1 = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    last_seen = ''
    for col in range(first_col, last_col):
        if col in building_by_col:
            last_seen = building_by_col[col]
            continue
        value = str(row1[col] or '').strip() if col < len(row1) else ''
        if value:
            building_by_col[col] = value
            last_seen = value
        elif last_seen:
            building_by_col[col] = last_seen
    return building_by_col


def _build_floor_cols(worksheet, headers, email_id_col):
    """
    Every non-blank header after the Internal Department Email column is a
    (building, floor) pair whose cells hold that floor's responsible email
    DIRECTLY - a filled cell means the checkpoint applies to that floor with
    that owner; a blank cell means it doesn't apply. There is no separate
    applicability-marker column to cross-reference by name (which is fragile
    - e.g. an abbreviated "G floor" header not matching "Ground Floor" in the
    email block would silently drop that floor).
    """
    building_by_col = _row1_building_by_col(worksheet, email_id_col + 1, len(headers))
    floor_cols = []
    for i in range(email_id_col + 1, len(headers)):
        floor_name = headers[i]
        if not floor_name:
            continue
        floor_cols.append({
            'building': building_by_col.get(i) or 'Default Building',
            'name': floor_name,
            'email_col': i,
        })
    return floor_cols


def _build_floor_assignment(row, fc):
    ecol = fc['email_col']
    responsible = str(row[ecol] or '').strip() if len(row) > ecol else ''
    return {
        'building': fc['building'],
        'floor': fc['name'],
        'applicable': bool(responsible),
        'responsible_email': responsible,
    }


def _parse_checklist_data_rows(data_rows, floor_cols, email_id_col):
    conditions = []
    current_condition_obj = None
    for row in data_rows:
        if not row or not any(cell is not None for cell in row):
            continue
        condition_name = str(row[0] or '').strip() if row else ''
        checkpoint_text = str(row[1] or '').strip() if len(row) > 1 else ''
        if condition_name:
            current_condition_obj = {'name': condition_name, 'checkpoints': []}
            conditions.append(current_condition_obj)
        if not current_condition_obj or not checkpoint_text:
            continue
        internal_email = str(row[email_id_col] or '').strip() if len(row) > email_id_col else ''
        current_condition_obj['checkpoints'].append({
            'name': checkpoint_text,
            'internal_email': internal_email,
            'floor_assignments': [_build_floor_assignment(row, fc) for fc in floor_cols],
        })
    return conditions


def _parse_checklist_workbook(workbook, sheet_name):
    """Parse a checklist workbook sheet and return structured data.

    Requires a workbook loaded WITHOUT read_only=True - merged-cell ranges
    (needed to resolve building groups from row 1) are unavailable on
    openpyxl's read-only worksheets.
    """
    ws = workbook[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError('File has too few rows to parse.')
    title_row, header_row, data_rows = rows[0], rows[1], rows[2:]
    headers = [str(h or '').strip() for h in header_row]
    email_id_col = _find_email_separator_col(headers)
    floor_cols = _build_floor_cols(ws, headers, email_id_col)
    if not floor_cols:
        raise ValueError(
            'No floor columns found after the Internal Department Email column.'
        )
    conditions = _parse_checklist_data_rows(data_rows, floor_cols, email_id_col)

    buildings = []
    seen_buildings = set()
    for fc in floor_cols:
        if fc['building'] not in seen_buildings:
            seen_buildings.add(fc['building'])
            buildings.append(fc['building'])

    return {
        'title': str(title_row[0] or '').strip() if title_row else '',
        'buildings': buildings,
        'floor_cols': floor_cols,
        'conditions': conditions,
    }


@login_required
def import_checklist_page(request):
    """Excel checklist import page."""
    if not is_warehouse_admin(request.user):
        return render(request, _WI_ACCESS_DENIED_TEMPLATE)
    return render(request, 'warehouse_inspection/import_checklist.html', {
        'page_title': 'Import Checklist from Excel',
        'is_admin': True,
    })


@login_required
@require_http_methods(['POST'])
def import_checklist_preview(request):
    """Parse an uploaded Excel file and return a JSON preview."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    try:
        import openpyxl
    except ImportError:
        return JsonResponse({'error': 'openpyxl is required. Run: pip install openpyxl'}, status=500)

    uploaded_file = request.FILES.get('file')
    if not uploaded_file:
        return JsonResponse({'error': 'No file uploaded.'}, status=400)

    sheet_name = request.POST.get('sheet', 'Sheet1')

    try:
        # NOT read_only - merged-cell ranges (needed to resolve building
        # groups from row 1) are unavailable on openpyxl's read-only worksheets.
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
    except Exception as exc:
        return JsonResponse({'error': f'Could not read workbook: {exc}'}, status=400)

    if sheet_name not in wb.sheetnames:
        return JsonResponse({
            'error': f"Sheet '{sheet_name}' not found.",
            'available_sheets': wb.sheetnames,
        }, status=400)

    try:
        parsed = _parse_checklist_workbook(wb, sheet_name)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)

    conditions = parsed['conditions']
    total_checkpoints = sum(len(c['checkpoints']) for c in conditions)
    total_assignments = sum(
        sum(1 for fa in cp['floor_assignments'] if fa['applicable'])
        for c in conditions
        for cp in c['checkpoints']
    )

    return JsonResponse({
        'sheets': wb.sheetnames,
        'title': parsed['title'],
        'buildings': parsed['buildings'],
        'floor_cols': [
            {'building': fc['building'], 'name': fc['name']} for fc in parsed['floor_cols']
        ],
        'conditions': conditions,
        'stats': {
            'buildings': len(parsed['buildings']),
            'conditions': len(conditions),
            'checkpoints': total_checkpoints,
            'floor_assignments': total_assignments,
            'floors': len(parsed['floor_cols']),
        },
    })


def _import_floor_assignments(checkpoint, floor_assignments, building_name, floor_by_key, counters):
    """Wire up only THIS checkpoint's own building's floors - `floor_assignments`
    carries every building's cells for this checkpoint row, since one sheet
    row spans every building's columns."""
    for fa in floor_assignments:
        if not fa['applicable'] or fa['building'] != building_name:
            continue
        floor = floor_by_key.get((fa['building'], fa['floor']))
        if not floor:
            continue
        _, assign_created = WarehouseInspectionFloorAssignment.objects.update_or_create(
            checkpoint=checkpoint,
            floor=floor,
            defaults={'is_applicable': True, 'responsible_email': fa['responsible_email'] or None},
        )
        if assign_created:
            counters['assignments'] += 1


def _import_conditions_for_building(building, conditions, floor_by_key, created_by, counters):
    """
    Create/update this building's own ConditionCategory + Checkpoint rows.
    A checklist that conceptually applies to several buildings still needs
    one independent set of rows per building - ConditionCategory and
    Checkpoint each FK to exactly one building, so the same condition/
    checkpoint names are intentionally duplicated across buildings, same as
    if each building had been imported from its own single-building sheet.
    """
    cp_order = {}
    for condition_data in conditions:
        condition, cond_created = ConditionCategory.objects.update_or_create(
            building=building,
            condition_category_name=condition_data['name'],
            defaults={'entity_type': 'admin', 'created_by': created_by, 'is_deleted': False},
        )
        if cond_created:
            counters['conditions'] += 1
        cid = condition.condition_category_id
        cp_order.setdefault(cid, 0)
        for cp_data in condition_data['checkpoints']:
            cp_order[cid] += 1
            checkpoint, cp_created = WarehouseInspectionCheckpoint.objects.update_or_create(
                condition_category=condition,
                checkpoint_name=cp_data['name'],
                defaults={
                    'checkpoint_order': cp_order[cid],
                    'is_active': True,
                    'created_by': created_by,
                    'internal_department_email': cp_data['internal_email'] or None,
                },
            )
            if cp_created:
                counters['checkpoints'] += 1
            _import_floor_assignments(
                checkpoint, cp_data['floor_assignments'], building.building_name, floor_by_key, counters,
            )


def _persist_checklist_import(floor_cols, conditions, created_by):
    """
    Buildings, and their floors, are discovered from the sheet itself
    (floor_cols carries each column's resolved building name) rather than
    supplied manually - one upload can populate several buildings at once.
    """
    counters = {'buildings': 0, 'floors': 0, 'conditions': 0, 'checkpoints': 0, 'assignments': 0}
    with transaction.atomic():
        building_by_name = {}
        floor_by_key = {}
        floor_order_by_building = {}

        for fc in floor_cols:
            building_name = fc['building']
            if building_name not in building_by_name:
                building, building_created = WarehouseInspectionBuilding.objects.update_or_create(
                    building_name=building_name,
                    defaults={'is_active': True},
                )
                building_by_name[building_name] = building
                if building_created:
                    counters['buildings'] += 1

            key = (building_name, fc['name'])
            if key not in floor_by_key:
                floor_order_by_building[building_name] = floor_order_by_building.get(building_name, 0) + 1
                floor, floor_created = WarehouseInspectionFloor.objects.update_or_create(
                    building=building_by_name[building_name],
                    floor_name=fc['name'],
                    defaults={'floor_order': floor_order_by_building[building_name], 'is_active': True},
                )
                floor_by_key[key] = floor
                if floor_created:
                    counters['floors'] += 1

        for building in building_by_name.values():
            _import_conditions_for_building(building, conditions, floor_by_key, created_by, counters)
    return counters


@login_required
@require_http_methods(['POST'])
def import_checklist_confirm(request):
    """Import a parsed Excel file into the warehouse inspection database."""
    if not is_warehouse_admin(request.user):
        return JsonResponse({'error': _WI_ERR_ACCESS_DENIED}, status=403)

    try:
        import openpyxl
    except ImportError:
        return JsonResponse({'error': 'openpyxl is required.'}, status=500)

    uploaded_file = request.FILES.get('file')
    if not uploaded_file:
        return JsonResponse({'error': 'No file uploaded.'}, status=400)

    sheet_name = request.POST.get('sheet', 'Sheet1')
    created_by = request.user.email or 'scm.digitalsupport@merillife.com'

    try:
        # NOT read_only - merged-cell ranges (needed to resolve building
        # groups from row 1) are unavailable on openpyxl's read-only worksheets.
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
    except Exception as exc:
        return JsonResponse({'error': f'Could not read workbook: {exc}'}, status=400)

    if sheet_name not in wb.sheetnames:
        return JsonResponse({'error': f"Sheet '{sheet_name}' not found."}, status=400)

    try:
        parsed = _parse_checklist_workbook(wb, sheet_name)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)

    floor_cols = parsed['floor_cols']
    conditions = parsed['conditions']

    try:
        counters = _persist_checklist_import(floor_cols, conditions, created_by)
    except Exception as exc:  # noqa: BLE001 — import errors surfaced to caller
        return JsonResponse({'error': f'Import failed: {exc}'}, status=500)

    buildings_label = ', '.join(f'"{name}"' for name in parsed['buildings'])
    return JsonResponse({
        'success': True,
        'buildings': parsed['buildings'],
        'counters': counters,
        'message': (
            f"Successfully imported: {counters['buildings']} building(s), "
            f"{counters['floors']} floor(s), {counters['conditions']} condition(s), "
            f"{counters['checkpoints']} checkpoint(s), "
            f"{counters['assignments']} floor assignment(s) across {buildings_label}."
        ),
    })


# ────────────────────────────────────────────────────────────────
# Scheduled daily-email trigger (n8n / external scheduler)
# ────────────────────────────────────────────────────────────────

class WarehouseInspectionDailyEmailAPIView(APIView):
    """
    External scheduler entry point for the daily warehouse inspection emails.

    Runs exactly the same work as the
    `send_warehouse_inspection_reminders` management command:
      1. tickets.send_ticket_reminder_emails()      - "still No" reminders
      2. services.send_daily_warehouse_inspection_report() - DGM daily report

    Exists so the schedule can live in n8n (which runs on a server) instead of
    a Windows Scheduled Task on one person's laptop - that task only fires when
    that machine happens to be awake at 09:00.

    Usage:
      curl -X POST https://gpt.e-meril.in/warehouse-inspection/api/run-daily-emails/ \
           -H 'X-WAREHOUSE-RUN-SECRET: <secret>' \
           -H 'Content-Type: application/json' \
           -d '{"mode": "all"}'

    Body (all optional):
      mode         all (default) | reminders | report
      report_date  YYYY-MM-DD - which day the report covers (default: yesterday)

    Protected by WAREHOUSE_INSPECTION_RUN_SECRET (configured in .env). Emails
    are sent synchronously, so allow a generous client timeout.
    """

    permission_classes = [AllowAny]
    authentication_classes = []

    def post(self, request, *args, **kwargs):
        provided_secret = (
            request.headers.get('X-WAREHOUSE-RUN-SECRET')
            or (request.data.get('run_secret') if hasattr(request.data, 'get') else None)
        )
        configured_secret = getattr(settings, 'WAREHOUSE_INSPECTION_RUN_SECRET', '')
        if not configured_secret or provided_secret != configured_secret:
            logger.warning('Warehouse inspection daily-email request rejected - invalid secret')
            return Response({'detail': 'Invalid credentials'}, status=status.HTTP_403_FORBIDDEN)

        mode = str((request.data.get('mode') if hasattr(request.data, 'get') else None) or 'all').lower()
        if mode not in {'all', 'reminders', 'report'}:
            return Response(
                {'detail': 'mode must be one of: all, reminders, report.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        raw_date = (request.data.get('report_date') if hasattr(request.data, 'get') else None) or ''
        report_date = None
        if raw_date:
            try:
                report_date = datetime.strptime(str(raw_date).strip(), '%Y-%m-%d').date()
            except ValueError:
                return Response(
                    {'detail': 'report_date must be in YYYY-MM-DD format.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        from .services import default_report_date, send_daily_warehouse_inspection_report
        from .tickets import send_ticket_reminder_emails

        empty = {'recipients': [], 'cc': [], 'sent_count': 0, 'failed_count': 0, 'errors': []}
        reminders = dict(empty, skipped=True)
        report = dict(empty, skipped=True)

        if mode in {'all', 'reminders'}:
            reminders = dict(send_ticket_reminder_emails(), skipped=False)
        if mode in {'all', 'report'}:
            resolved_date = report_date or default_report_date()
            report = dict(send_daily_warehouse_inspection_report(resolved_date), skipped=False)
            report['report_date'] = resolved_date.isoformat()

        # A single flag so the n8n side needs no expression gymnastics to
        # decide whether to raise an alert.
        ok = not (
            reminders.get('failed_count') or reminders.get('errors')
            or report.get('failed_count') or report.get('errors')
        )
        payload = {
            'ok': ok,
            'mode': mode,
            'ran_at': timezone.localtime().isoformat(),
            'reminders': reminders,
            'report': report,
        }
        if not ok:
            logger.error('Warehouse inspection daily emails completed with problems: %s', payload)
        return Response(payload, status=status.HTTP_200_OK)
