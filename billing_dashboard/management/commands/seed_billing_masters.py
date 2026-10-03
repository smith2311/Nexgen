"""
Seed the billing enrichment masters from the stakeholder's "Lookup" sheet
(the multi-block sheet shipped inside the Endo Export Billing Tracker workbook).

    python manage.py seed_billing_masters --path "/path/to/Endo export billing tracker.xlsx"
    python manage.py seed_billing_masters --path <file> --sheet Lookup --dry-run

The Lookup sheet lays each master out in a fixed column range (headers in row 1):

    A-D  Soldto-Name → Country, Subs-flag, Region      → SoldToPartyMap
    O-P  Material Code (exact) → LOB                    → LobMap (MATERIAL_CODE)
    R-S  Country → Division ("Push Report Division")    → CountryRegionMap.push_division

Idempotent (``update_or_create``) — safe to re-run when the Lookup is refreshed.
"""
import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from billing_dashboard.models import CountryRegionMap, LobMap, SoldToPartyMap
from billing_dashboard.services.enrichment import normalize_party_name

# Domestic markers seen in the Lookup's Soldto-Name block.
_DOMESTIC_FLAGS = {"india"}
_DOMESTIC_REGIONS = {"domestic", "mumbai office", "india"}
# Header/label values to skip per block.
_PARTY_HEADERS = {"soldto-name", "sold-to name", ""}
_LOB_HEADERS = {"material group", "material code", ""}
_DIV_HEADERS = {"country", ""}


def _s(v):
    return "" if v is None else str(v).strip()


def parse_lookup_workbook(path, sheet="Lookup"):
    """Parse the multi-block Lookup sheet → (party, lob_codes, divisions).

    * party      : {soldto_name_key: {...SoldToPartyMap fields...}}
    * lob_codes  : {material_code: lob}
    * divisions  : {country_name: push_division}
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CommandError(f"Could not open '{path}': {exc}")
    if sheet not in wb.sheetnames:
        raise CommandError(f"Sheet '{sheet}' not found. Available: {wb.sheetnames}")
    ws = wb[sheet]

    party, lob_codes, divisions = {}, {}, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        def cell(i):
            return _s(row[i]) if i < len(row) else ""

        # A-D → Soldto-Name master
        name = cell(0)
        if name and name.lower() not in _PARTY_HEADERS:
            key = normalize_party_name(name)
            if key:
                subs_flag = cell(2)
                region = cell(3)
                party[key] = dict(
                    soldto_name=name, country=cell(1) or None, region=region or None,
                    subs_flag=subs_flag or None,
                    is_subsidiary=subs_flag.lower() == "subsidiary",
                    is_domestic=(subs_flag.lower() in _DOMESTIC_FLAGS
                                 or region.lower() in _DOMESTIC_REGIONS),
                )

        # O-P → Material Code → LOB
        mcode, mlob = cell(14), cell(15)
        if mcode and mcode.lower() not in _LOB_HEADERS and mlob and mlob.lower() != "lob":
            lob_codes[mcode] = mlob

        # R-S → Country → Push Report Division
        dcountry, ddiv = cell(17), cell(18)
        if dcountry and dcountry.lower() not in _DIV_HEADERS and ddiv and ddiv.lower() != "division":
            divisions[dcountry] = ddiv

    wb.close()
    return party, lob_codes, divisions


class Command(BaseCommand):
    help = "Seed SoldToPartyMap / LOB (material code) / Country→Division from the Lookup sheet."

    def add_arguments(self, parser):
        parser.add_argument("--path", required=True, help="Path to the tracker .xlsx containing the Lookup sheet.")
        parser.add_argument("--sheet", default="Lookup", help="Sheet name (default: Lookup).")
        parser.add_argument("--dry-run", action="store_true", help="Parse & report counts without writing.")

    def handle(self, *args, **opts):
        party, lob_codes, divisions = parse_lookup_workbook(opts["path"], opts["sheet"])
        self.stdout.write(
            f"Parsed Lookup: {len(party)} sold-to parties, {len(lob_codes)} material→LOB, "
            f"{len(divisions)} country→division."
        )
        if opts["dry_run"]:
            self.stdout.write(self.style.WARNING("Dry run — nothing written."))
            return

        with transaction.atomic():
            p_new = self._seed_party(party)
            l_new = self._seed_lob(lob_codes)
            d_upd = self._seed_divisions(divisions)

        self.stdout.write(self.style.SUCCESS(
            f"Seeded: SoldToPartyMap +{p_new}/{len(party)}, "
            f"LOB(material-code) +{l_new}/{len(lob_codes)}, "
            f"Country→Division {d_upd}/{len(divisions)} upserted."
        ))

    def _seed_party(self, party):
        created = 0
        for key, data in party.items():
            _, was_new = SoldToPartyMap.objects.update_or_create(
                soldto_name_key=key,
                defaults={**data, "is_active": True},
            )
            created += int(was_new)
        return created

    def _seed_lob(self, lob_codes):
        created = 0
        for code, lob in lob_codes.items():
            _, was_new = LobMap.objects.update_or_create(
                key_type=LobMap.KeyType.MATERIAL_CODE, key_value=code,
                defaults={"lob": lob, "is_active": True},
            )
            created += int(was_new)
        return created

    def _seed_divisions(self, divisions):
        upserted = 0
        for country, div in divisions.items():
            CountryRegionMap.objects.update_or_create(
                country_name=country,
                defaults={"push_division": div},
            )
            upserted += 1
        return upserted
