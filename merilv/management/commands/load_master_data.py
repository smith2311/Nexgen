"""Load the SAP-derived "Master Data.xlsx" workbook straight into Meril-V's masters.

Reads every relevant sheet and upserts:
  - MHPL / MLSIPL Material Master  -> MaterialMaster (per entity)
  - 9100 / 3100 Customer Master    -> CustomerMaster (per entity, by sales org)
  - Billing and Sales Type         -> BillingTypeMaster, SalesTypeMaster
  - Plant-900 / Plant - 300        -> PlantMaster
  - Storage-900 / Storage - 300    -> StorageMaster
  - User Details                   -> Django users + a Meril-V role group

Not imported: the "Headers" sheet — that tab documents YVL06/MB52 export
column names for reference, it is not master data.

Idempotent — safe to re-run; existing rows are updated in place, not duplicated.

Usage (from inside the running container):
    python manage.py load_master_data
    python manage.py load_master_data --file "/app/Master Data.xlsx"
    python manage.py load_master_data --skip users types
"""

import openpyxl
from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.core.management.base import BaseCommand, CommandError

from merilv.models import (
    BillingTypeMaster,
    CustomerMaster,
    Entity,
    MaterialMaster,
    PlantMaster,
    SalesTypeMaster,
    StorageMaster,
)
from merilv.services import ingest

SECTIONS = ("materials", "customers", "types", "plants", "storages", "users")

#: The sheet's free-text "Role" column only ever says Admin or User (plus one
#: instructional non-row for an unregistered Picker) — mapped onto the app's
#: three real tiers. "User" is treated as day-to-day operational access
#: (Picker), not read-only (Viewer); adjust per-person afterward from
#: Master Data -> Roles & Access if that is wrong for someone.
ROLE_MAP = {"admin": "Meril-V Admin", "user": "Meril-V Picker", "picker": "Meril-V Picker", "viewer": "Meril-V Viewer"}
ALL_ROLE_GROUPS = ("Meril-V Viewer", "Meril-V Picker", "Meril-V Admin")

#: Model field max_length limits — external SAP data is untrusted, so long
#: values are truncated defensively rather than crashing a bulk insert.
_LIMITS = {
    "scan_code": 100, "material_code": 100, "material_description": 400,
    "gtin": 20, "material_group": 100,
    "customer_code": 50, "customer_name": 300, "company_code": 20, "plant": 20,
    "courier_partner": 150, "courier_number": 100, "city": 120, "state": 120,
    "country": 120, "pincode": 20, "gstin": 30, "email": 254, "mobile": 40,
    "contact_person": 150,
}


def _clip(value, field):
    limit = _LIMITS.get(field)
    return value[:limit] if limit and value else value


def _model_limits(model):
    """max_length for every char field on a model, keyed by field name.

    Derived from the model rather than hand-listed in _LIMITS: the YMM061 set
    is 33 fields wide and growing, and a hand-maintained table would silently
    stop clipping the moment someone adds a column — which surfaces as a bulk
    insert crashing on one long SAP description, thousands of rows in.
    """
    return {
        f.name: f.max_length
        for f in model._meta.get_fields()
        if getattr(f, "max_length", None)
    }


class Command(BaseCommand):
    help = 'Load materials, customers, plants, storages, billing/sales types and users from "Master Data.xlsx".'

    def add_arguments(self, parser):
        parser.add_argument("--file", default="Master Data.xlsx", help="Path to the workbook.")
        parser.add_argument(
            "--skip", nargs="*", choices=SECTIONS, default=[],
            help="Sections to skip: " + ", ".join(SECTIONS),
        )

    def handle(self, *args, **options):
        path = options["file"]
        skip = set(options["skip"])

        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {path}")
        except Exception as exc:
            raise CommandError(f"Could not open {path}: {exc}")

        call_command("seed_merilv")
        entities = {e.code: e for e in Entity.objects.filter(code__in=("MHPL", "MLSIPL"))}
        if "MHPL" not in entities or "MLSIPL" not in entities:
            raise CommandError("Entities MHPL/MLSIPL are missing even after seed_merilv — aborting.")

        try:
            if "materials" not in skip:
                self._load_materials(wb, entities)
            if "customers" not in skip:
                self._load_customers(wb, entities)
            if "types" not in skip:
                self._load_billing_sales_types(wb)
            if "plants" not in skip:
                self._load_plants(wb)
            if "storages" not in skip:
                self._load_storages(wb)
            if "users" not in skip:
                self._load_users(wb)
        finally:
            wb.close()

        self.stdout.write(self.style.SUCCESS("Master data load complete."))

    # ── helpers ──────────────────────────────────────────────────────────
    def _sheet_grid(self, wb, name):
        rows = list(wb[name].iter_rows(values_only=True))
        headers = ["" if c is None else str(c).strip() for c in rows[0]]
        return headers, rows[1:]

    def _bulk_upsert(self, model, objs, unique_fields, update_fields, batch_size=2000):
        for i in range(0, len(objs), batch_size):
            model.objects.bulk_create(
                objs[i : i + batch_size],
                update_conflicts=True,
                unique_fields=unique_fields,
                update_fields=update_fields,
                batch_size=batch_size,
            )

    # ── materials ────────────────────────────────────────────────────────
    def _load_materials(self, wb, entities):
        plan = [("MHPL Material Master", entities["MHPL"]), ("MLSIPL Material Master", entities["MLSIPL"])]
        total = 0
        for sheet_name, entity in plan:
            headers, rows = self._sheet_grid(wb, sheet_name)
            result = ingest.parse_material_master_grid(headers, rows)

            # Every YMM061 field the parser resolved, not just the five this
            # used to write. The parser has always produced the whole set —
            # resolved by exact SAP header, which matters because the MHPL and
            # MLSIPL sheets carry the same columns in different orders — but
            # the loader constructed each row from a hardcoded handful, so
            # PH Levels, EAN/UPC, material type and division arrived parsed and
            # were dropped on the floor. They existed only inside raw_data,
            # where nothing can filter, group or display them; that is the
            # "Master Data population failure — missing Material Group, PH
            # Levels" report.
            ymm_fields = [
                field for _header, field in MaterialMaster.YMM061_COLUMNS
                if field not in ("scan_code",)
            ]
            limits = _model_limits(MaterialMaster)

            objs = []
            for r in result.records:
                obj = MaterialMaster(
                    entity=entity,
                    scan_code=_clip(r["scan_code"], "scan_code"),
                    material_code=_clip(r["material_code"], "material_code"),
                    material_description=_clip(r["material_description"], "material_description"),
                    gtin=_clip(r["gtin"], "gtin"),
                    material_group=_clip(r["material_group"], "material_group"),
                    raw_data=r["raw_data"],
                )
                for field in ymm_fields:
                    if field not in r:
                        continue
                    value = r[field]
                    if field in MaterialMaster.YMM061_BOOLEAN_FIELDS:
                        # SAP writes "X", the model stores a boolean. Slicing a
                        # letter to max_length would leave the string "X" going
                        # into a BooleanField.
                        setattr(obj, field, MaterialMaster.coerce_flag(value))
                        continue
                    limit = limits.get(field)
                    setattr(obj, field, value[:limit] if limit and value else value)
                # bulk_create bypasses save(), so the derived flags are applied
                # by hand — otherwise every imported row would come in with
                # serial_managed False regardless of its profile.
                obj.apply_derived_flags()
                objs.append(obj)

            self._bulk_upsert(
                MaterialMaster, objs,
                unique_fields=["entity", "scan_code"],
                update_fields=sorted(
                    {"material_code", "material_description", "gtin", "material_group", "raw_data",
                     # serial_managed is derived rather than mapped from a
                     # header, so it is not in ymm_fields and has to be named.
                     "serial_managed", *ymm_fields} - {"scan_code"}
                ),
            )
            total += len(objs)
            note = f" ({result.skipped} skipped)" if result.skipped else ""
            self.stdout.write(f"  materials [{entity.code}]: {len(objs):,} rows from '{sheet_name}'{note}")
        self.stdout.write(self.style.SUCCESS(
            f"Materials: {total:,} rows upserted. MaterialMaster now has {MaterialMaster.objects.count():,} rows."
        ))

    # ── customers ────────────────────────────────────────────────────────
    def _load_customers(self, wb, entities):
        # Sales org 9100 -> Meril Healthcare Pvt. Ltd. (MHPL); 3100 -> Meril Life
        # Sciences India Pvt. Ltd. (MLSIPL) — matches Entity.company_name exactly
        # (see merilv/services/entities.py) and the two sales orgs share zero
        # customer codes, so there is no cross-entity collision risk.
        plan = [("9100 - Customer Master", entities["MHPL"]), ("3100 - Customer Master", entities["MLSIPL"])]
        total = 0
        for sheet_name, entity in plan:
            headers, rows = self._sheet_grid(wb, sheet_name)
            result = ingest.parse_customer_master_grid(headers, rows)

            # A key only for a column the sheet actually had. The parser is
            # already careful about this (see parse_customer_master_grid), but
            # reading every field with `.get(key, "")` here threw that away:
            # an absent column arrived as "", and listing it in update_fields
            # wrote that "" over every existing row.
            #
            # That is not hypothetical. "Master Data.xlsx" is the SAP customer
            # extract this command loads on EVERY container start, and neither
            # of its customer sheets has a courier column — so each restart
            # blanked courier_partner/courier_number on all 7,000-odd
            # customers, and the outbound console's courier auto-fetch went
            # dead until somebody re-imported the Courier/Customer master by
            # hand. An absent COLUMN means no opinion; an empty CELL in a
            # column that IS present still clears the field, because that is
            # somebody deliberately doing so.
            optional = [
                "customer_name", "company_code", "plant", "courier_partner", "courier_number",
                "address", "city", "state", "country", "pincode", "gstin", "email",
                "mobile", "contact_person",
            ]
            present = [f for f in optional if any(f in r for r in result.records)]

            objs = []
            for r in result.records:
                values = {f: r[f] for f in present if f in r}
                for field, value in values.items():
                    if field != "address":          # address is the one free-text column
                        values[field] = _clip(value, field)
                objs.append(CustomerMaster(
                    customer_code=_clip(r["customer_code"], "customer_code"),
                    entity=entity,
                    raw_data=r["raw_data"],
                    **values,
                ))
            self._bulk_upsert(
                CustomerMaster, objs,
                unique_fields=["customer_code"],
                update_fields=[*present, "entity", "raw_data"],
            )
            total += len(objs)
            note = f" ({result.skipped} skipped)" if result.skipped else ""
            self.stdout.write(f"  customers [{entity.code}]: {len(objs):,} rows from '{sheet_name}'{note}")
        self.stdout.write(self.style.SUCCESS(
            f"Customers: {total:,} rows upserted. CustomerMaster now has {CustomerMaster.objects.count():,} rows."
        ))

    # ── billing & sales types ────────────────────────────────────────────
    def _load_billing_sales_types(self, wb):
        """The sheet is one table with four independent code/description blocks
        (MHPL sales, MLSIPL sales, MLSIPL billing, MHPL billing) under a two-row
        header. BillingTypeMaster/SalesTypeMaster are global, not per-entity, so
        the two blocks of each type are unioned; where a code means something
        different for MHPL vs MLSIPL, MHPL's description is kept and every such
        collision is reported so it isn't a silent decision.
        """
        rows = list(wb["Billing and Sales Type"].iter_rows(values_only=True))

        def block(code_i, desc_i):
            out = {}
            for r in rows[2:]:
                code = r[code_i]
                if code is None or str(code).strip() == "":
                    continue
                out[str(code).strip()] = str(r[desc_i]).strip() if r[desc_i] else ""
            return out

        mhpl_sales, mlsipl_sales = block(0, 1), block(3, 4)
        mlsipl_bill, mhpl_bill = block(6, 7), block(9, 10)

        sales = {**mlsipl_sales, **mhpl_sales}  # MHPL wins on conflict
        billing = {**mlsipl_bill, **mhpl_bill}  # MHPL wins on conflict

        conflicts = sorted(
            {k for k in (mhpl_sales.keys() & mlsipl_sales.keys()) if mhpl_sales[k] != mlsipl_sales[k]}
            | {k for k in (mhpl_bill.keys() & mlsipl_bill.keys()) if mhpl_bill[k] != mlsipl_bill[k]}
        )
        if conflicts:
            self.stdout.write(self.style.WARNING(
                f"  {len(conflicts)} sales/billing type code(s) mean something different for MHPL vs "
                f"MLSIPL (e.g. {conflicts[:5]}) — this master is shared across entities, so MHPL's "
                "description was kept for each. Check Master Data if a specific code matters."
            ))

        for code, desc in sales.items():
            SalesTypeMaster.objects.update_or_create(sales_type=code, defaults={"sales_description": desc})
        for code, desc in billing.items():
            BillingTypeMaster.objects.update_or_create(billing_type=code, defaults={"billing_description": desc})
        self.stdout.write(self.style.SUCCESS(f"Sales types: {len(sales)}. Billing types: {len(billing)}."))

    # ── plants ───────────────────────────────────────────────────────────
    def _load_plants(self, wb):
        # No Company/Organisation code exists anywhere in the source data (only
        # the plant's own code and a legal-entity name repeated across rows), so
        # PlantMaster.organization is deliberately left unset rather than guessed.
        plan = ["Plant-900", "Plant - 300"]
        total = 0
        for sheet_name in plan:
            headers, rows = self._sheet_grid(wb, sheet_name)
            idx = {h: i for i, h in enumerate(headers)}
            code_i, term_i = idx.get("Plnt"), idx.get("Search Term 1")
            name2_i, name_i = idx.get("Name 2"), idx.get("Name")
            for r in rows:
                code = r[code_i] if code_i is not None else None
                if code in (None, ""):
                    continue
                code = str(code).strip()
                label = (
                    (str(r[term_i]).strip() if term_i is not None and r[term_i] else "")
                    or (str(r[name2_i]).strip() if name2_i is not None and r[name2_i] else "")
                    or (str(r[name_i]).strip() if name_i is not None and r[name_i] else "")
                    or f"Plant {code}"
                )
                PlantMaster.objects.update_or_create(plant_code=code, defaults={"plant_name": label[:200]})
                total += 1
        self.stdout.write(self.style.SUCCESS(
            f"Plants: {total} rows upserted from Plant-900 / Plant - 300 "
            "(no Company/Organisation link — not present in the source data)."
        ))

    # ── storages ─────────────────────────────────────────────────────────
    def _load_storages(self, wb):
        plan = ["Storage-900", "Storage - 300"]
        total = auto_plants = 0
        auto_codes = []
        for sheet_name in plan:
            headers, rows = self._sheet_grid(wb, sheet_name)
            idx = {h: i for i, h in enumerate(headers)}
            plant_i = idx.get("Plnt", idx.get("Plant"))
            code_i = idx.get("SLoc", idx.get("Storage"))
            desc_i = idx.get("Description")
            for r in rows:
                plant_code = r[plant_i] if plant_i is not None else None
                code = r[code_i] if code_i is not None else None
                if plant_code in (None, "") or code in (None, ""):
                    continue
                plant_code = str(plant_code).strip()
                plant, created = PlantMaster.objects.get_or_create(
                    plant_code=plant_code, defaults={"plant_name": f"Unmapped Plant {plant_code}"},
                )
                if created:
                    auto_plants += 1
                    auto_codes.append(plant_code)
                StorageMaster.objects.update_or_create(
                    storage_code=str(code).strip(), plant=plant,
                    defaults={"storage_name": str(r[desc_i]).strip() if desc_i is not None and r[desc_i] else ""},
                )
                total += 1
        msg = f"Storage locations: {total} rows upserted."
        if auto_plants:
            msg += (
                f" {auto_plants} plant code(s) referenced by storage but absent from the Plant sheets "
                f"were auto-created as placeholders ({', '.join(auto_codes)}) — rename them in "
                "Master Data -> Organisation."
            )
        self.stdout.write(self.style.SUCCESS(msg))

    # ── users ────────────────────────────────────────────────────────────
    def _load_users(self, wb):
        headers, rows = self._sheet_grid(wb, "User Details")
        idx = {h: i for i, h in enumerate(headers)}
        name_i, email_i, role_i = idx["Name"], idx["Email"], idx["Role"]
        created = updated = skipped = 0
        for r in rows:
            name = str(r[name_i]).strip() if r[name_i] else ""
            email = str(r[email_i]).strip() if r[email_i] else ""
            role_raw = str(r[role_i]).strip().lower() if r[role_i] else ""
            # One row is an instruction ("Picker need To register"), not a real
            # person — no name and no real email, so it carries nothing to load.
            if not name or "@" not in email:
                skipped += 1
                continue
            group_name = ROLE_MAP.get(role_raw)
            if not group_name:
                self.stdout.write(self.style.WARNING(f"  unknown role '{role_raw}' for {email} — skipped."))
                skipped += 1
                continue

            parts = name.replace("Mr.", "").replace("Ms.", "").replace("Mrs.", "").strip().split(None, 1)
            first, last = (parts[0], parts[1]) if len(parts) > 1 else (parts[0] if parts else "", "")

            user = User.objects.filter(email__iexact=email).first()
            if user is None:
                user = User.objects.create(
                    username=email, email=email, first_name=first, last_name=last, is_active=True,
                )
                user.set_unusable_password()
                user.save()
                created += 1
            else:
                updated += 1

            group, _ = Group.objects.get_or_create(name=group_name)
            other_roles = Group.objects.filter(name__in=[g for g in ALL_ROLE_GROUPS if g != group_name])
            user.groups.remove(*other_roles)
            user.groups.add(group)

        self.stdout.write(self.style.SUCCESS(
            f"Users: {created} created, {updated} already existed (role refreshed), {skipped} skipped."
        ))
