import json
from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.db.models import Sum, Count

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ui.decorators import uses_components
from .models import SubLogisticEntry


# ── Role helpers ───────────────────────────────────────────────────────────────

def _sub_role(user):
    if not user.is_authenticated:
        return None
    if user.is_superuser or user.groups.filter(name="Sub-Manager").exists():
        return "manager"
    if user.groups.filter(name="Sub-Member").exists():
        return "member"
    return None


def _sub_login_required(view_fn):
    from functools import wraps

    @wraps(view_fn)
    def _inner(request, *args, **kwargs):
        if not request.user.is_authenticated:
            from django.conf import settings
            return redirect(f"{settings.LOGIN_URL}?next={request.path}")
        if _sub_role(request.user) is None:
            from django.core.exceptions import PermissionDenied
            raise PermissionDenied
        return view_fn(request, *args, **kwargs)
    return _inner


def _sub_require_manager(view_fn):
    from functools import wraps

    @wraps(view_fn)
    def _inner(request, *args, **kwargs):
        if not request.user.is_authenticated:
            from django.conf import settings
            return redirect(f"{settings.LOGIN_URL}?next={request.path}")
        if _sub_role(request.user) != "manager":
            from django.core.exceptions import PermissionDenied
            raise PermissionDenied
        return view_fn(request, *args, **kwargs)
    return _inner


# ── Excel parsing helpers ──────────────────────────────────────────────────────

def _parse_dec(val):
    if val is None or str(val).strip() in ("", "-", "—"):
        return None
    try:
        return Decimal(str(val).replace(",", "").strip())
    except InvalidOperation:
        return None


def _parse_date(val):
    if val is None:
        return None
    if hasattr(val, "date"):          # openpyxl datetime cell
        return val.date()
    s = str(val).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_upload(fileobj):
    """
    Read an .xlsx file and return (rows, errors).
    rows  = list of dicts with keys: country, total_sales, ob_logistics_cost, operational_cost, period_date
    errors = list of str
    """
    wb = load_workbook(fileobj, data_only=True)
    ws = wb.active
    rows   = []
    errors = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(c for c in row):
            continue
        country     = str(row[0]).strip() if row[0] else None
        total_sales = _parse_dec(row[1]) if len(row) > 1 else None
        ob_cost     = _parse_dec(row[2]) if len(row) > 2 else None
        ops_cost    = _parse_dec(row[3]) if len(row) > 3 else None
        pdate       = _parse_date(row[4]) if len(row) > 4 else None

        row_errs = []
        if not country:
            row_errs.append("Country is required")
        if total_sales is None:
            row_errs.append("Total Sales is required")
        if pdate is None:
            row_errs.append("Date is required (DD-MM-YYYY)")

        rows.append({
            "row":              i,
            "country":          country or "",
            "total_sales":      str(total_sales) if total_sales is not None else "",
            "ob_logistics_cost": str(ob_cost)     if ob_cost     is not None else "",
            "operational_cost": str(ops_cost)    if ops_cost    is not None else "",
            "period_date":      pdate.strftime("%Y-%m-%d") if pdate else "",
            "errors":           row_errs,
            "valid":            len(row_errs) == 0,
        })
        if row_errs:
            errors.append(f"Row {i}: {'; '.join(row_errs)}")

    return rows, errors


# ── Views ──────────────────────────────────────────────────────────────────────

@_sub_login_required
@uses_components("data_table", "master_page")
def home(request):
    return render(request, "subsidiary_logistic/home.html", {
        "sub_role": _sub_role(request.user),
    })


@_sub_login_required
@uses_components("data_table", "master_page")
def dashboard(request):
    role = _sub_role(request.user)
    qs   = SubLogisticEntry.objects.filter(is_deleted=False)

    # Date filter
    date_filter = request.GET.get("date", "").strip()
    active_date = None
    if date_filter:
        try:
            active_date = datetime.strptime(date_filter, "%Y-%m-%d").date()
            qs = qs.filter(period_date=active_date)
        except ValueError:
            pass

    entries = list(qs.order_by("period_date", "country"))

    # Aggregates for header stats
    totals = qs.aggregate(
        t_sales = Sum("total_sales"),
        t_ob    = Sum("ob_logistics_cost"),
        t_ops   = Sum("operational_cost"),
        t_count = Count("country", distinct=True),
    )

    # Available dates for filter dropdown
    dates = (
        SubLogisticEntry.objects
        .filter(is_deleted=False)
        .values_list("period_date", flat=True)
        .distinct()
        .order_by("-period_date")
    )

    return render(request, "subsidiary_logistic/dashboard.html", {
        "entries":     entries,
        "totals":      totals,
        "dates":       dates,
        "active_date": active_date,
        "date_filter": date_filter,
        "sub_role":    role,
    })


@_sub_require_manager
@uses_components("data_table", "master_page")
def upload(request):
    if request.method == "POST":
        action = request.POST.get("action", "")

        if action == "save":
            rows_json = request.POST.get("rows_json", "[]")
            try:
                rows = json.loads(rows_json)
            except (json.JSONDecodeError, TypeError):
                messages.error(request, "Invalid data — please re-upload and preview again.")
                return redirect("sub:upload")

            saved = 0
            for r in rows:
                if not r.get("valid"):
                    continue
                try:
                    SubLogisticEntry.objects.create(
                        country           = r["country"],
                        total_sales       = Decimal(r["total_sales"]),
                        ob_logistics_cost = Decimal(r["ob_logistics_cost"]) if r.get("ob_logistics_cost") else None,
                        operational_cost  = Decimal(r["operational_cost"])  if r.get("operational_cost")  else None,
                        period_date       = datetime.strptime(r["period_date"], "%Y-%m-%d").date(),
                    )
                    saved += 1
                except Exception:
                    pass

            messages.success(request, f"{saved} record{'s' if saved != 1 else ''} saved successfully.")
            return redirect("sub:dashboard")

    return render(request, "subsidiary_logistic/upload.html", {
        "sub_role": _sub_role(request.user),
    })


@_sub_require_manager
def upload_preview(request):
    """AJAX endpoint: receive .xlsx, parse, return JSON rows for preview."""
    if request.method != "POST":
        return JsonResponse({"error": "POST required"}, status=405)

    f = request.FILES.get("excel_file")
    if not f:
        return JsonResponse({"error": "No file provided"}, status=400)
    if not f.name.lower().endswith(".xlsx"):
        return JsonResponse({"error": "Only .xlsx files are accepted"}, status=400)

    try:
        rows, errors = _parse_upload(f)
    except Exception as e:
        return JsonResponse({"error": f"Could not parse file: {e}"}, status=400)

    valid_count = sum(1 for r in rows if r["valid"])
    return JsonResponse({
        "rows":        rows,
        "total":       len(rows),
        "valid":       valid_count,
        "error_count": len(rows) - valid_count,
        "errors":      errors,
    })


@_sub_login_required
def template_download(request):
    """Download a styled .xlsx template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sub Logistics Data"

    headers = [
        "Country",
        "Total Sales (In Cr.)",
        "OB Logistics Cost (In Cr.)",
        "Operational Cost (In Cr.)",
        "Date (DD-MM-YYYY)",
    ]

    hdr_fill = PatternFill("solid", fgColor="0B1F3A")
    hdr_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    align_c  = Alignment(horizontal="center", vertical="center")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill      = hdr_fill
        cell.font      = hdr_font
        cell.alignment = align_c

    # Sample rows
    samples = [
        ("GMBH",           "46.88", "1.65", "0.73", "01-04-2026"),
        ("Brazil",         "29.15", "1.76", "0.71", "01-04-2026"),
        ("Russia",         "79.25", "",     "",      "01-04-2026"),
        ("South Africa",   "2.80",  "0.22", "0.15",  "01-04-2026"),
        ("Turkey",         "62.43", "0.42", "0.28",  "01-04-2026"),
    ]
    for r, row in enumerate(samples, 2):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)

    col_widths = [20, 22, 26, 24, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws.row_dimensions[1].height = 22

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="sub_logistics_template.xlsx"'
    wb.save(response)
    return response
